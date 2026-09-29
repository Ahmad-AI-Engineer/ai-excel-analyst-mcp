"""Check an upstream MCP chart in the saved workbook and place it safely."""

from __future__ import annotations

import os
import re
import tempfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import get_column_letter, quote_sheetname, range_boundaries

CHART_SHEET = "Charts"
DATA_SHEET = "_ChartData"


def _title(chart):
    try:
        return "".join(
            run.t
            for paragraph in chart.title.tx.rich.p
            for run in paragraph.r
        )
    except (AttributeError, TypeError):
        return ""


def _formula(reference):
    if reference is None:
        return None
    for kind in ("numRef", "strRef"):
        value = getattr(reference, kind, None)
        if value is not None:
            return value.f
    return None


def _rect(chart):
    anchor = chart.anchor
    if not hasattr(anchor, "_from"):
        return None
    start = anchor._from
    if hasattr(anchor, "to"):
        return (start.col + 1, start.row + 1, anchor.to.col + 1, anchor.to.row + 1)
    return (start.col + 1, start.row + 1, start.col + 11, start.row + 17)


def _overlap(a, b):
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def _reference_bounds(formula):
    if not formula or "!" not in formula:
        return None
    name, cells = formula.rsplit("!", 1)
    try:
        return name.strip("'").replace("''", "'"), range_boundaries(cells.replace("$", ""))
    except ValueError:
        return None


def _move_legacy_chart_content(book):
    """Keep existing Charts cells intact in a hidden archive, with live chart refs."""
    display = book[CHART_SHEET] if CHART_SHEET in book else book.create_sheet(CHART_SHEET)
    display.sheet_view.showGridLines = False
    if not any(cell.value is not None for cell in display._cells.values()):
        return display, None

    # Renaming a sheet does not rewrite ordinary formula references in openpyxl.
    # Refuse a migration that would silently change the meaning of user formulas.
    for sheet in book:
        for cell in sheet._cells.values():
            if cell.data_type != "f" or not isinstance(cell.value, str):
                continue
            if any(token.type == "OPERAND" and token.subtype == "RANGE"
                   and _reference_bounds(token.value)
                   and _reference_bounds(token.value)[0].casefold() == CHART_SHEET.casefold()
                   for token in Tokenizer(cell.value).items):
                raise ValueError("Charts contains referenced cells; cannot safely make it chart-only")

    archive = "_ChartArchive"
    index = 2
    while archive in book:
        archive = f"_ChartArchive{index}"
        index += 1
    original_index = book.sheetnames.index(CHART_SHEET)
    display.title = archive
    display.sheet_state = "hidden"
    clean = book.create_sheet(CHART_SHEET, original_index)
    clean.sheet_view.showGridLines = False

    for sheet in book:
        for chart in sheet._charts:
            for series in chart.series:
                for reference in (series.cat, series.val, series.tx):
                    for kind in ("numRef", "strRef"):
                        item = getattr(reference, kind, None) if reference is not None else None
                        formula = getattr(item, "f", None)
                        bounds = _reference_bounds(formula)
                        if bounds and bounds[0].casefold() == CHART_SHEET.casefold():
                            item.f = f"{quote_sheetname(archive)}!{formula.rsplit('!', 1)[1]}"
    for row, chart in enumerate(list(display._charts), start=0):
        chart.visible_cells_only = False
        chart.anchor = TwoCellAnchor(
            _from=AnchorMarker(col=1, row=1 + row * 18),
            to=AnchorMarker(col=11, row=17 + row * 18),
        )
        clean.add_chart(chart)
    display._charts.clear()
    return clean, archive


def _aggregate_source(book, source, value_name, category_name, chart_sheet, question):
    """Find the raw sheet with both requested headers and its complete data range."""
    preferred_name = None
    preferred_bounds = None
    if "!" in source:
        name, cells = source.rsplit("!", 1)
        preferred_name = name.strip("'").replace("''", "'")
        try:
            preferred_bounds = range_boundaries(cells.replace("$", ""))
        except ValueError:
            pass
    candidates = []
    for worksheet in book:
        if worksheet.title in {chart_sheet, DATA_SHEET} and worksheet.title != preferred_name:
            continue
        populated = [cell for cell in worksheet._cells.values() if cell.value is not None]
        if not populated:
            continue
        header_row = min(cell.row for cell in populated)
        first_col = min(cell.column for cell in populated if cell.row == header_row)
        last_col = first_col
        while last_col < worksheet.max_column and worksheet.cell(header_row, last_col + 1).value is not None:
            last_col += 1
        if worksheet.title == preferred_name and preferred_bounds:
            preferred_left, preferred_top, preferred_right, _ = preferred_bounds
            preferred_headers = {str(worksheet.cell(preferred_top, col).value).strip().lower(): col
                                 for col in range(preferred_left, preferred_right + 1)}
            if value_name in preferred_headers and category_name in preferred_headers:
                first_col, header_row, last_col = preferred_left, preferred_top, preferred_right
        headers = {}
        for col in range(first_col, last_col + 1):
            headers.setdefault(str(worksheet.cell(header_row, col).value).strip().lower(), col)
        if value_name in headers and category_name in headers:
            last_row = max((cell.row for cell in populated
                            if cell.column in {headers[value_name], headers[category_name]}
                            and cell.row >= header_row), default=header_row)
            candidates.append((worksheet, (first_col, header_row, last_col, last_row), headers))
    hinted = [item for item in candidates if re.search(
        rf"\bfrom\s+{re.escape(item[0].title)}\b", question, re.I)]
    preferred = [item for item in candidates if item[0].title == preferred_name]
    if len(hinted) == 1:
        selected = hinted[0]
    elif len(preferred) == 1:
        selected = preferred[0]
    elif len(candidates) == 1:
        selected = candidates[0]
    elif not candidates:
        raise ValueError(f"No source sheet has both {category_name} and {value_name} columns")
    else:
        raise ValueError(f"Several sheets have {category_name} and {value_name}; specify the source sheet")
    worksheet, bounds, headers = selected
    left, top, right, bottom = bounds
    reference = f"{quote_sheetname(worksheet.title)}!{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"
    return worksheet, bounds, headers, reference


def verify_chart(filepath: str, sheet_name: str, data_range: str, title: str,
                 chart_type: str = "", question: str = "", legacy_sheet: str = "") -> dict:
    """Verify actual series, replace a matching old chart and save a clear anchor.

    The range must be a two-column category/value summary on the chart sheet.
    This is intentionally post-processing a saved MCP result, not another MCP tool.
    """
    if "!" in data_range:
        source_sheet, data_range = data_range.rsplit("!", 1)
        if source_sheet.strip("'") != sheet_name:
            raise ValueError("Chart data must be prepared on its destination sheet")
    left, top, right, bottom = range_boundaries(data_range.replace("$", ""))
    if right != left + 1 or bottom <= top:
        raise ValueError("Chart source must be a two-column range with headers and data")
    book = load_workbook(filepath)
    cached = load_workbook(filepath, data_only=True)
    try:
        def source_value(cell):
            if cell.data_type != "f":
                return cell.value
            value = cached[cell.parent.title][cell.coordinate].value
            if value is None:
                raise ValueError(f"Formula at {cell.parent.title}!{cell.coordinate} has no calculated value")
            return value

        sheet = book[sheet_name]
        if not sheet._charts:
            raise ValueError("MCP reported success but no chart was saved")
        matches = [item for item in sheet._charts if _title(item) == title] if title else sheet._charts
        if not matches:
            raise ValueError("MCP reported success but the requested chart was not saved")
        chart = next((item for item in reversed(matches) if not isinstance(item.anchor, TwoCellAnchor)), matches[-1])
        if len(chart.series) != 1:
            raise ValueError("Saved chart must have exactly one value series")
        expected_cat = f"'{sheet_name}'!${sheet.cell(top + 1, left).column_letter}${top + 1}:${sheet.cell(bottom, left).column_letter}${bottom}"
        expected_val = f"'{sheet_name}'!${sheet.cell(top + 1, right).column_letter}${top + 1}:${sheet.cell(bottom, right).column_letter}${bottom}"
        series = chart.series[0]
        actual_cat = _formula(series.cat)
        actual_val = _formula(series.val)
        def normal(formula):
            return (formula or "").replace("'", "").replace("$", "").lower()
        if normal(actual_cat) != normal(expected_cat) or normal(actual_val) != normal(expected_val):
            raise ValueError(f"Saved chart series mismatch: category={actual_cat}, value={actual_val}")
        if not sheet.cell(top, left).value or not sheet.cell(top, right).value:
            raise ValueError("Chart summary headers are missing")
        if not any(sheet.cell(row, left).value is not None for row in range(top + 1, bottom + 1)):
            raise ValueError("Chart categories are empty")
        if not any(sheet.cell(row, right).value is not None for row in range(top + 1, bottom + 1)):
            raise ValueError("Chart values are empty")
        if title and _title(chart) != title:
            raise ValueError("Saved chart title does not match the request")

        # For grouped SUM charts, compare every plotted value with the source
        # rows. Metadata is written by the agent through ordinary MCP tools.
        blocks = []
        for cell in list(sheet._cells.values()):
            if str(cell.value).strip().lower() != "title" or sheet.cell(cell.row, cell.column + 1).value != title:
                continue
            block = {}
            for row in range(cell.row, min(cell.row + 4, sheet.max_row + 1)):
                key = str(sheet.cell(row, cell.column).value or "").strip().lower()
                if key in {"title", "source", "request", "type"}:
                    block[key] = sheet.cell(row, cell.column + 1).value
            blocks.append((block, (cell.column, cell.row, cell.column + 2, cell.row + 4),
                           abs(cell.column - left) + abs(cell.row - top)))
        metadata, metadata_rect, _ = min(blocks, key=lambda item: item[2]) if blocks else ({}, None, 0)
        if set(metadata) != {"title", "source", "request", "type"}:
            raise ValueError("Chart is missing its Title / Source / Request / Type block")
        # The saved chart is authoritative. Metadata is descriptive and can use
        # a different name (for example "column") for the same chart family.
        requested_type = chart_type.lower().strip()
        if requested_type in {"bar", "column"} and not isinstance(chart, BarChart):
            raise ValueError("Saved chart is not a bar chart")
        if requested_type == "pie" and not isinstance(chart, PieChart):
            raise ValueError("Saved chart is not a pie chart")
        request = str(metadata.get("request") or "")
        aggregate = re.search(r"sum\(([^)]+)\)\s+by\s+(.+)", request, re.IGNORECASE)
        if not aggregate:
            total_title = re.fullmatch(r"\s*total\s+(.+?)\s+by\s+(.+?)\s*", title, re.I)
            if total_title:
                request = f"SUM({total_title.group(1).strip()}) by {total_title.group(2).strip()}"
                for row in range(metadata_rect[1], metadata_rect[3]):
                    if str(sheet.cell(row, metadata_rect[0]).value).strip().lower() == "request":
                        sheet.cell(row, metadata_rect[0] + 1).value = request
                        break
                metadata["request"] = request
                aggregate = re.search(r"sum\(([^)]+)\)\s+by\s+(.+)", request, re.IGNORECASE)
        if ("total" in title.lower() or "sum" in title.lower()) and not aggregate:
            raise ValueError("Grouped total chart is missing a SUM(value) by category request block")
        if aggregate:
            source = str(metadata.get("source") or "")
            value_name, category_name = aggregate.group(1).strip().lower(), aggregate.group(2).strip().lower()
            source_sheet, source_bounds, headers, corrected_source = _aggregate_source(
                book, source, value_name, category_name, sheet_name, question)
            source_left, source_top, source_right, source_bottom = source_bounds
            if source != corrected_source:
                for row in range(metadata_rect[1], metadata_rect[3]):
                    if str(sheet.cell(row, metadata_rect[0]).value).strip().lower() == "source":
                        sheet.cell(row, metadata_rect[0] + 1).value = corrected_source
                        break
                metadata["source"] = corrected_source
            totals = defaultdict(Decimal)
            for source_row in range(source_top + 1, source_bottom + 1):
                category = source_value(source_sheet.cell(source_row, headers[category_name]))
                raw_value = source_value(source_sheet.cell(source_row, headers[value_name]))
                if category is None or raw_value is None:
                    continue
                try:
                    totals[str(category)] += Decimal(str(raw_value))
                except InvalidOperation as exc:
                    raise ValueError("SUM source includes a nonnumeric value") from exc
            plotted = {}
            for summary_row in range(top + 1, bottom + 1):
                category = source_value(sheet.cell(summary_row, left))
                value = source_value(sheet.cell(summary_row, right))
                try:
                    plotted[str(category)] = Decimal(str(value))
                except InvalidOperation as exc:
                    raise ValueError("Chart summary must contain saved numeric totals") from exc
            if plotted.keys() != totals.keys() or any(abs(plotted[key] - totals[key]) > Decimal("0.000001") for key in totals):
                raise ValueError("Chart summary does not equal SUM by category over the full source range")

        display, archived_sheet = _move_legacy_chart_content(book) if sheet_name == DATA_SHEET else (sheet, None)
        if sheet_name == DATA_SHEET:
            sheet._charts.remove(chart)
            display.add_chart(chart)
            chart.visible_cells_only = False
            sheet.sheet_state = "hidden"
        # Place the complete 10-by-16-cell chart clear of populated cells and charts.
        last_col = max((cell.column for cell in display._cells.values() if cell.value is not None), default=0)
        last_row = max((cell.row for cell in display._cells.values() if cell.value is not None), default=0)
        other_charts = []
        original_source = str(metadata["source"])
        source_ref = _reference_bounds(original_source)
        if source_ref is None:
            raise ValueError("Chart Source must name a worksheet and cell range")
        if archived_sheet and source_ref[0].casefold() == CHART_SHEET.casefold():
            original_source = f"{quote_sheetname(archived_sheet)}!{original_source.rsplit('!', 1)[1]}"
            source_ref = _reference_bounds(original_source)
            sheet.cell(metadata_rect[1] + 1, metadata_rect[0] + 1).value = original_source
            metadata["source"] = original_source
        for old in list(display._charts):
            if old is chart:
                continue
            old_cat = _reference_bounds(_formula(old.series[0].cat)) if old.series else None
            old_val = _reference_bounds(_formula(old.series[0].val)) if old.series else None
            same_summary = (old.series and normal(_formula(old.series[0].cat)) == normal(expected_cat)
                            and normal(_formula(old.series[0].val)) == normal(expected_val))
            same_raw_source = bool(
                source_ref and old_cat and old_val
                and old_cat[0] == source_ref[0] and old_val[0] == source_ref[0]
                and all(source_ref[1][0] <= ref[1][0] <= ref[1][2] <= source_ref[1][2]
                        and source_ref[1][1] <= ref[1][1] <= ref[1][3] <= source_ref[1][3]
                        for ref in (old_cat, old_val))
            )
            old_header = (old_cat[1][0], old_cat[1][1] - 1) if old_cat else None
            old_metadata_sheet = book[old_cat[0]] if old_cat and old_cat[0] in book else None
            same_metadata_source = bool(old_header and old_metadata_sheet and any(
                str(old_metadata_sheet.cell(cell.row + 1, cell.column).value).strip().lower() == "source"
                and _reference_bounds(str(old_metadata_sheet.cell(cell.row + 1, cell.column + 1).value)) == source_ref
                and str(old_metadata_sheet.cell(cell.row, cell.column + 1).value) == _title(old)
                and abs(cell.column - old_header[0]) + abs(cell.row - old_header[1]) <= 8
                for cell in list(old_metadata_sheet._cells.values())
                if str(cell.value).strip().lower() == "title"))
            same_identity = same_summary or same_raw_source or same_metadata_source
            updating = bool(re.search(r"\b(update|fix|replace)\b", question, re.I))
            if same_identity and (_title(old) == title or updating):
                display._charts = [item for item in display._charts if item is not old]
            else:
                bounds = _rect(old)
                if bounds:
                    other_charts.append(bounds)
        summary_rect = (left, top, right + 1, bottom + 1)
        if display is sheet and any(_overlap(summary_rect, bounds) or _overlap(metadata_rect, bounds) for bounds in other_charts):
            raise ValueError("Chart summary or metadata overlaps an existing chart")
        if legacy_sheet and legacy_sheet not in {sheet_name, display.title} and legacy_sheet in book:
            legacy = book[legacy_sheet]
            for old in list(legacy._charts):
                if not title or _title(old) != title or not old.series:
                    continue
                old_cat = _reference_bounds(_formula(old.series[0].cat))
                old_val = _reference_bounds(_formula(old.series[0].val))
                same_raw_source = bool(
                    old_cat and old_val and old_cat[0] == source_ref[0] and old_val[0] == source_ref[0]
                    and all(source_ref[1][0] <= ref[1][0] <= ref[1][2] <= source_ref[1][2]
                            and source_ref[1][1] <= ref[1][1] <= ref[1][3] <= source_ref[1][3]
                            for ref in (old_cat, old_val))
                )
                header = (old_cat[1][0], old_cat[1][1] - 1) if old_cat and old_cat[0] == legacy_sheet else None
                same_metadata_source = bool(header and any(
                    str(legacy.cell(cell.row, cell.column + 1).value) == title
                    and str(legacy.cell(cell.row + 1, cell.column).value).strip().lower() == "source"
                    and _reference_bounds(str(legacy.cell(cell.row + 1, cell.column + 1).value)) == source_ref
                    and abs(cell.column - header[0]) + abs(cell.row - header[1]) <= 8
                    for cell in list(legacy._cells.values()) if str(cell.value).strip().lower() == "title"))
                if same_raw_source or same_metadata_source:
                    legacy._charts = [item for item in legacy._charts if item is not old]
        right_side = last_col + 11 <= 16384
        col, row = (last_col + 2, 2) if right_side else (2, last_row + 3)
        candidate = (col, row, col + 10, row + 16)
        while any(_overlap(candidate, bounds) for bounds in other_charts):
            row += 18
            candidate = (col, row, col + 10, row + 16)
        if candidate[3] > 1048576:
            raise ValueError("No empty chart area remains on the destination sheet")
        if (right_side and candidate[0] <= last_col) or (not right_side and candidate[1] <= last_row):
            raise ValueError("Chart overlaps populated cells")
        chart.anchor = TwoCellAnchor(
            _from=AnchorMarker(col=col - 1, row=row - 1),
            to=AnchorMarker(col=col + 9, row=row + 15),
        )
        folder = os.path.dirname(filepath)
        fd, temporary = tempfile.mkstemp(suffix=".xlsx", dir=folder)
        os.close(fd)
        try:
            book.save(temporary)
            check = load_workbook(temporary)
            saved = next((item for item in check[display.title]._charts
                          if _title(item) == title and item.series
                          and normal(_formula(item.series[0].cat)) == normal(expected_cat)
                          and normal(_formula(item.series[0].val)) == normal(expected_val)), None)
            if saved is None:
                raise ValueError("Chart disappeared during save")
            if _rect(saved) != candidate or normal(_formula(saved.series[0].cat)) != normal(expected_cat) or normal(_formula(saved.series[0].val)) != normal(expected_val):
                raise ValueError("Chart failed saved-workbook verification")
            if sheet_name == DATA_SHEET and (check[DATA_SHEET].sheet_state != "hidden"
                                             or any(cell.value is not None for cell in check[CHART_SHEET]._cells.values())
                                             or saved.visible_cells_only is not False):
                raise ValueError("Charts sheet is not chart-only after saving")
            check.close()
            os.replace(temporary, filepath)
        finally:
            if os.path.exists(temporary):
                os.remove(temporary)
        return {"verified": True, "sheet": display.title, "title": title, "category_ref": actual_cat,
                "value_ref": actual_val, "bounds": candidate}
    finally:
        book.close()
        cached.close()
