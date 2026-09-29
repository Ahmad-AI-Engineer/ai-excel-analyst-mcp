"""Calculate XLSX formulas and update only their cached worksheet values."""

from __future__ import annotations

import html
import math
import os
import posixpath
import re
import tempfile
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.datetime import to_excel


_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
_WHOLE_COLUMNS = re.compile(
    r"^(?:(?P<sheet>'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.]*)!)?"
    r"(?P<first>\$?[A-Z]{1,3}):(?P<last>\$?[A-Z]{1,3})$",
    re.IGNORECASE,
)
_CELL = re.compile(rb"(?P<open><c(?=[\s>])[^>]*>)(?P<body>.*?)(?P<close></c>)", re.DOTALL)
_COORDINATE = re.compile(rb'\br="([A-Z]+[1-9][0-9]*)"')
_FORMULA = re.compile(rb"<f(?:\s[^>]*)?(?:/>|>.*?</f>)", re.DOTALL)
_VALUE = re.compile(rb"<v(?:\s[^>]*)?(?:/>|>.*?</v>)", re.DOTALL)
_TYPE = re.compile(rb'\s+t="[^"]*"')
_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_OFFICE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _bound_columns(formula: str, current_sheet: str, used_rows: dict[str, int]) -> str:
    """Limit whole-column lookups in a disposable model to populated rows."""
    tokens = Tokenizer(formula).items
    changed = False
    parts = []
    for token in tokens:
        value = token.value
        if token.type == "OPERAND" and token.subtype == "RANGE":
            match = _WHOLE_COLUMNS.fullmatch(value)
            if match:
                sheet = match.group("sheet")
                name = sheet.strip("'").replace("''", "'") if sheet else current_sheet
                matching_rows = next((rows for title, rows in used_rows.items()
                                      if title.casefold() == name.casefold()), None)
                if matching_rows is None:
                    raise ValueError(f"Formula references missing sheet {name}")
                last = max(matching_rows, 1)
                prefix = f"{sheet}!" if sheet else ""
                value = f"{prefix}{match.group('first')}1:{match.group('last')}{last}"
                changed = True
        parts.append(value)
    return "=" + "".join(parts) if changed else formula


def _formula_cells(filepath: str):
    book = load_workbook(filepath, data_only=False)
    try:
        formulas = {}
        used_rows = {}
        bounded = {}
        for sheet in book.worksheets:
            populated = [cell.row for cell in sheet._cells.values() if cell.value is not None]
            used_rows[sheet.title] = max(populated, default=1)
        for sheet in book.worksheets:
            for cell in sheet._cells.values():
                if cell.data_type != "f":
                    continue
                if not isinstance(cell.value, str):
                    raise ValueError(f"Unsupported formula type at {sheet.title}!{cell.coordinate}")
                key = (sheet.title, cell.coordinate)
                formulas[key] = cell.value
                converted = _bound_columns(cell.value, sheet.title, used_rows)
                if converted != cell.value:
                    cell.value = converted
                    bounded[key] = converted
        return book, formulas, bounded
    except Exception:
        book.close()
        raise


def _scalar(value, location: str, epoch):
    value = getattr(value, "value", value)
    if hasattr(value, "size"):
        if value.size != 1:
            raise ValueError(f"Array formula at {location} cannot be cached as one cell")
        value = value.item()
    if value is None or str(value) in _ERRORS:
        raise ValueError(f"Formula at {location} calculated to {value or 'no value'}")
    if isinstance(value, (datetime, date)):
        value = to_excel(value, epoch=epoch)
    if isinstance(value, Decimal):
        value = float(value)
    if not isinstance(value, (bool, int, float, str)):
        raise ValueError(f"Unsupported formula result at {location}: {type(value).__name__}")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"Formula at {location} calculated to a non-finite number")
    if isinstance(value, float) and value.is_integer() and abs(value) <= 2**53:
        value = int(value)
    return value


def _encoded_value(value):
    if isinstance(value, bool):
        return b"b", b"1" if value else b"0"
    if isinstance(value, (int, float)):
        return b"n", str(value).encode("ascii")
    return b"str", html.escape(value, quote=False).encode("utf-8")


def _worksheet_paths(archive: zipfile.ZipFile) -> dict[str, str]:
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
    paths = {}
    for sheet in workbook.findall(f".//{{{_MAIN}}}sheet"):
        target = targets[sheet.attrib[f"{{{_OFFICE_REL}}}id"]]
        path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        paths[sheet.attrib["name"]] = path
    return paths


def _patch_xml(xml: bytes, values: dict[str, object]) -> bytes:
    pending = {key.encode("ascii"): value for key, value in values.items()}

    def replace(match):
        opened, body, closed = match.group("open", "body", "close")
        coordinate = _COORDINATE.search(opened)
        if not coordinate or coordinate.group(1) not in pending:
            return match.group(0)
        key = coordinate.group(1)
        without_old_value = _VALUE.sub(b"", body)
        formula = _FORMULA.search(without_old_value)
        if not formula:
            raise ValueError(f"Saved cell {key.decode()} has no formula")
        kind, encoded = _encoded_value(pending.pop(key))
        opened = _TYPE.sub(b"", opened)
        opened = opened[:-1] + b' t="' + kind + b'">'
        cached = b"<v>" + encoded + b"</v>"
        body = without_old_value[:formula.end()] + cached + without_old_value[formula.end():]
        return opened + body + closed

    updated = _CELL.sub(replace, xml)
    if pending:
        raise ValueError("Formula cache cells missing from saved worksheet: " + ", ".join(key.decode() for key in pending))
    return updated


def _verify(filepath: str, formulas: dict, values: dict) -> None:
    with zipfile.ZipFile(filepath) as archive:
        if archive.testzip() is not None:
            raise ValueError("Recalculated workbook failed ZIP validation")
    formula_book = load_workbook(filepath, read_only=True, data_only=False)
    value_book = load_workbook(filepath, read_only=True, data_only=True)
    try:
        for (sheet, coordinate), original in formulas.items():
            location = f"{sheet}!{coordinate}"
            if formula_book[sheet][coordinate].value != original:
                raise ValueError(f"Formula changed while caching {location}")
            actual, expected = value_book[sheet][coordinate].value, values[(sheet, coordinate)]
            if isinstance(actual, (datetime, date)):
                actual = to_excel(actual, epoch=value_book.epoch)
            if expected == "" and actual is None:
                continue
            if isinstance(expected, (int, float)) and not isinstance(expected, bool):
                if not isinstance(actual, (int, float)) or not math.isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-10):
                    raise ValueError(f"Cached value was not saved for {location}")
            elif actual != expected:
                raise ValueError(f"Cached value was not saved for {location}")
    finally:
        formula_book.close()
        value_book.close()


def recalculate_formula_caches(filepath: str) -> bool:
    """Recalculate formula cells; preserve every non-worksheet ZIP member."""
    if not filepath.lower().endswith(".xlsx"):
        return False
    book, formulas, bounded = _formula_cells(filepath)
    try:
        if not formulas:
            return False
        with tempfile.TemporaryDirectory() as calculation_dir:
            model_path = filepath
            if bounded:
                model_path = os.path.join(calculation_dir, Path(filepath).name)
                book.save(model_path)
            book.close()
            import formulas as formula_engine

            model = formula_engine.ExcelModel().loads(model_path).finish()
            solution = {str(key).upper(): value for key, value in model.calculate().items()}
            calculated = {}
            for (sheet, coordinate) in formulas:
                key = f"'[{Path(model_path).name}]{sheet.replace(chr(39), chr(39) * 2)}'!{coordinate}".upper()
                if key not in solution:
                    raise ValueError(f"Formula at {sheet}!{coordinate} could not be evaluated")
                calculated[(sheet, coordinate)] = _scalar(solution[key], f"{sheet}!{coordinate}", book.epoch)
        fd, temporary = tempfile.mkstemp(prefix="~$", suffix=".xlsx", dir=os.path.dirname(filepath))
        os.close(fd)
        try:
            with zipfile.ZipFile(filepath) as source, zipfile.ZipFile(temporary, "w") as output:
                paths = _worksheet_paths(source)
                by_path = {}
                for (sheet, coordinate), value in calculated.items():
                    by_path.setdefault(paths[sheet], {})[coordinate] = value
                for member in source.infolist():
                    content = source.read(member.filename)
                    if member.filename in by_path:
                        content = _patch_xml(content, by_path[member.filename])
                    output.writestr(member, content)
            _verify(temporary, formulas, calculated)
            os.replace(temporary, filepath)
        finally:
            if os.path.exists(temporary):
                os.remove(temporary)
        return True
    finally:
        book.close()
