"""Route MCP chart calls to a dedicated Charts worksheet."""

from __future__ import annotations

import os
import re
import tempfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from agents.chart_review import CHART_SHEET, DATA_SHEET, _aggregate_source


def _matching_metadata(sheet, title, left, top):
    matches = []
    for cell in list(sheet._cells.values()):
        if str(cell.value).strip().lower() != "title" or sheet.cell(cell.row, cell.column + 1).value != title:
            continue
        block = {}
        for row in range(cell.row, cell.row + 4):
            key = str(sheet.cell(row, cell.column).value or "").strip().lower()
            if key in {"title", "source", "request", "type"}:
                block[key] = sheet.cell(row, cell.column + 1).value
        matches.append((abs(cell.column - left) + abs(cell.row - top), block))
    return min(matches, key=lambda item: item[0])[1] if matches else {}


def _request_text(title: str, question: str) -> str:
    for value in (title, question):
        match = re.search(r"\btotal\s+(.+?)\s+by\s+(.+?)(?:\s+from\s+\w+|[.,;]|$)", value, re.I)
        if match:
            return f"SUM({match.group(1).strip()}) by {match.group(2).strip()}"
    return question


def _sum_fields(request: str, title: str, question: str):
    match = re.search(r"\bsum\(([^)]+)\)\s+by\s+([^.,;]+)", request, re.I)
    if match:
        category = re.split(r"\s+(?:from|in|on)\s+", match.group(2), maxsplit=1, flags=re.I)[0]
        return match.group(1).strip(), category.strip()
    generated = _request_text(title, question)
    match = re.fullmatch(r"SUM\(([^)]+)\) by (.+)", generated, re.I)
    if not match:
        return None
    category = re.split(r"\s+(?:from|in|on)\s+", match.group(2), maxsplit=1, flags=re.I)[0]
    return match.group(1).strip(), category.strip()


def prepare_chart_call(filepath: str, args: dict, question: str) -> dict:
    """Prepare chart data on a helper sheet and return routed MCP arguments.

    The caller keeps a workbook backup and restores it if MCP or saved-chart
    verification fails. Grouped SUM requests are calculated from the raw
    source, while other two-column ranges are copied. Source cells stay intact.
    """
    source_name = args["sheet_name"]
    data_range = args["data_range"]
    if "!" in data_range:
        source_name, data_range = data_range.rsplit("!", 1)
        source_name = source_name.strip("'").replace("''", "'")
    left, top, right, bottom = range_boundaries(data_range.replace("$", ""))

    book = load_workbook(filepath)
    cached = load_workbook(filepath, data_only=True)
    try:
        def source_value(cell):
            if cell.data_type != "f":
                return cell.value
            value = cached[cell.parent.title][cell.coordinate].value
            if value is None:
                raise ValueError(f"Formula at {cell.parent.title}!{cell.coordinate} has no calculated value")
            return value

        if source_name not in book:
            raise ValueError(f"Chart summary sheet {source_name} does not exist")
        source = book[source_name]
        title = args.get("title") or "Chart"
        metadata = _matching_metadata(source, title, left, top)
        fields = _sum_fields(str(metadata.get("request") or ""), title, question)
        if fields:
            value_name, category_name = fields
            raw_source = str(metadata.get("source") or f"'{source_name}'!{data_range}")
            raw_sheet, raw_bounds, headers, resolved_source = _aggregate_source(
                book, raw_source, value_name.lower(), category_name.lower(), CHART_SHEET, question)
            totals = defaultdict(Decimal)
            for row in range(raw_bounds[1] + 1, raw_bounds[3] + 1):
                category = source_value(raw_sheet.cell(row, headers[category_name.lower()]))
                amount = source_value(raw_sheet.cell(row, headers[value_name.lower()]))
                if category is None or amount is None:
                    continue
                try:
                    totals[str(category)] += Decimal(str(amount))
                except InvalidOperation as exc:
                    raise ValueError(f"{value_name} contains a nonnumeric value") from exc
            if not totals:
                raise ValueError("No numeric rows were found for the requested chart")
            values = [[category_name, f"Sum of {value_name}"]] + [
                [category, int(total) if total == int(total) else float(total)]
                for category, total in totals.items()]
            metadata = {"source": resolved_source,
                        "request": f"SUM({value_name}) by {category_name}"}
        else:
            if right != left + 1 or bottom <= top:
                raise ValueError("Prepare a two-column category/value range with headers before creating the chart")
            values = [[source_value(source.cell(row, col)) for col in (left, right)]
                      for row in range(top, bottom + 1)]
        if not all(values[0]) or not any(row[0] is not None and row[1] is not None for row in values[1:]):
            raise ValueError("Chart summary must have category and value headers and data")

        charts = book[CHART_SHEET] if CHART_SHEET in book else book.create_sheet(CHART_SHEET)
        charts.sheet_view.showGridLines = False
        destination = book[DATA_SHEET] if DATA_SHEET in book else book.create_sheet(DATA_SHEET)
        if any(cell.value is not None for cell in destination._cells.values()) and not any(
            str(cell.value).strip().lower() == "title" and cell.column == 4
            and str(destination.cell(cell.row + 1, 4).value).strip().lower() == "source"
            and str(destination.cell(cell.row + 2, 4).value).strip().lower() == "request"
            and str(destination.cell(cell.row + 3, 4).value).strip().lower() == "type"
            for cell in list(destination._cells.values())
        ):
            raise ValueError(f"{DATA_SHEET} already contains unrelated data")
        # The upstream MCP tool creates its chart on the data sheet. Keep this
        # sheet visible until the chart is moved and verified after the tool call.
        destination.sheet_state = "visible"
        used_row = max((cell.row for cell in destination._cells.values() if cell.value is not None), default=0)
        source_ref = metadata.get("source") or f"'{source_name}'!{data_range}"
        request_ref = metadata.get("request") or _request_text(title, question)
        chart_kind = args.get("chart_type", "")
        start = None
        previous_length = 0
        for cell in list(destination._cells.values()):
            row = cell.row
            if (cell.column == 4 and str(cell.value).strip().lower() == "title"
                    and destination.cell(row + 1, 5).value == source_ref
                    and destination.cell(row + 2, 5).value == request_ref
                    and str(destination.cell(row + 3, 5).value).casefold() == str(chart_kind).casefold()):
                length = 0
                while destination.cell(row + length, 1).value is not None or destination.cell(row + length, 2).value is not None:
                    length += 1
                if length >= len(values):
                    start, previous_length = row, length
                    break
        if start is None:
            start = used_row + 2 if used_row else 1
        for offset, row in enumerate(values):
            destination.cell(start + offset, 1, row[0])
            destination.cell(start + offset, 2, row[1])
        for row in range(start + len(values), start + previous_length):
            destination.cell(row, 1).value = None
            destination.cell(row, 2).value = None
        block = (
            ("Title", title),
            ("Source", source_ref),
            ("Request", request_ref),
            ("Type", chart_kind),
        )
        for offset, (key, value) in enumerate(block):
            destination.cell(start + offset, 4, key)
            destination.cell(start + offset, 5, value)
        fd, temporary = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(filepath))
        os.close(fd)
        try:
            book.save(temporary)
            os.replace(temporary, filepath)
        finally:
            if os.path.exists(temporary):
                os.remove(temporary)
        last = start + len(values) - 1
        return {**args, "sheet_name": DATA_SHEET,
                "data_range": f"A{start}:B{last}",
                "target_cell": f"G{start}", "title": title}
    finally:
        book.close()
        cached.close()
