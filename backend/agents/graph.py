import os
import asyncio
import tempfile
import shutil
import re
from typing import Optional, Any

from langgraph.checkpoint.memory import MemorySaver
from langgraph.store.memory import InMemoryStore
from langchain_mcp_adapters.client import MultiServerMCPClient
from langchain.agents import create_agent
from langchain.agents.middleware import wrap_model_call, wrap_tool_call
from langchain_openrouter import ChatOpenRouter
from langchain_core.messages import ToolMessage
from langgraph.errors import GraphBubbleUp
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from agents.chart_review import verify_chart
from agents.chart_prepare import CHART_SHEET, prepare_chart_call
from agents.chart_review import DATA_SHEET
from agents.formula_cache import recalculate_formula_caches
from workbook_sync import workbook_directory

UPLOAD_DIR = workbook_directory()

# Shared client instance
_mcp_client: Optional[MultiServerMCPClient] = None
_workbook_locks: dict[str, asyncio.Lock] = {}

# Shared memory instances for conversation continuity
_checkpointer = MemorySaver()
_store = InMemoryStore()


@wrap_model_call
async def retry_provider_errors(request, handler):
    """Retry only transient model-provider errors, without repeating Excel tools."""
    attempt = 0
    while True:
        try:
            return await handler(request)
        except GraphBubbleUp:
            raise
        except Exception as exc:
            error = str(exc).lower()
            if "provider_overloaded" in error:
                delays = (2.0, 4.0, 8.0)
            elif "provider_unavailable" in error:
                delays = (1.0, 2.0)
            else:
                raise
            if attempt >= len(delays):
                raise
            await asyncio.sleep(delays[attempt])
            attempt += 1


def pivot_destination_from_question(question: str) -> Optional[str]:
    if "pivot" not in question.lower():
        return None
    match = re.search(r"\b(?:sheet|worksheet)\s+(?:called|named)\s+[\"']?([^\"'.,;]+)", question, re.I)
    return match.group(1).strip() if match else None


def place_pivot_on_sheet(filepath: str, source_sheet: str, destination: str, original_sheets: set[str]) -> None:
    """Move the upstream `<source>_pivot` output to the requested sheet."""
    generated = f"{source_sheet}_pivot"
    if destination == generated:
        return
    book = load_workbook(filepath)
    try:
        if generated not in book:
            raise ValueError(f"Pivot tool did not save {generated}")
        if destination in original_sheets:
            if destination not in book:
                raise ValueError(f"The destination sheet {destination} disappeared")
            old = book[destination]
            new = book[generated]
            if len(old.tables) != 1 or old._charts:
                raise ValueError(f"The destination sheet {destination} contains other content")
            left, top, right, bottom = range_boundaries(next(iter(old.tables.values())).ref)
            new_right = new.max_column
            if (top != 1 or left != 1 or right != new_right
                    or [old.cell(1, col).value for col in range(1, right + 1)]
                    != [new.cell(1, col).value for col in range(1, new_right + 1)]
                    or any(cell.value is not None and not (left <= cell.column <= right and top <= cell.row <= bottom)
                           for cell in old._cells.values())):
                raise ValueError(f"The destination sheet {destination} contains other content")
        if destination in book:
            book.remove(book[destination])
        book[generated].title = destination
        fd, temporary = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(filepath))
        os.close(fd)
        try:
            book.save(temporary)
            check = load_workbook(temporary)
            try:
                if destination not in check or generated in check:
                    raise ValueError("Pivot destination was not saved")
            finally:
                check.close()
            os.replace(temporary, filepath)
        finally:
            if os.path.exists(temporary):
                os.remove(temporary)
    finally:
        book.close()

def get_mcp_client() -> MultiServerMCPClient:
    global _mcp_client
    if _mcp_client is None:
        stdio = {
            "transport": "stdio",
            "command": "uvx",
            "args": ["--from", "excel-mcp-server==0.1.8", "excel-mcp-server", "stdio"],
        }
        uv_environment = {key: os.environ[key] for key in ("UV_CACHE_DIR", "UV_TOOL_DIR", "UV_OFFLINE") if key in os.environ}
        if uv_environment:
            stdio["env"] = uv_environment
        _mcp_client = MultiServerMCPClient({
            "excel-service": stdio
        })
    return _mcp_client


async def aask_agent(
    question: str,
    file_path: Optional[str] = None,
    sheet_name: Optional[str] = None,
    thread_id: Optional[str] = "default_session",
) -> dict:
    """
    Executes the user's question using LangChain and official Excel MCP tools.
    1. Connects to the Excel MCP server via MultiServerMCPClient.
    2. Fetches the installed Excel MCP tools at runtime.
    3. Initializes OpenRouter chat model and LangChain agent with MemorySaver & InMemoryStore.
    4. Invokes the agent with conversation thread memory and returns the result.
    """
    filename = os.path.basename(file_path) if file_path else "a new workbook"
    session_id = thread_id or "default_session"
    
    # 1. Fetch MCP tools
    client = get_mcp_client()
    mcp_tools = await client.get_tools()
    operations = []
    errors = []
    pivot_destination = pivot_destination_from_question(question)
    pivot_source_match = re.search(r"\bfrom\s+([A-Za-z][A-Za-z0-9_-]*)", question, re.I) if pivot_destination else None
    pivot_source = pivot_source_match.group(1) if pivot_source_match else None
    original_sheets = set()
    if file_path and file_path.lower().endswith(".xlsx") and os.path.isfile(file_path):
        original_book = load_workbook(file_path)
        try:
            original_sheets = set(original_book.sheetnames)
        finally:
            original_book.close()
    write_tools = {"create_workbook", "create_worksheet", "write_data_to_excel", "apply_formula",
                   "format_range", "create_chart", "create_pivot_table", "create_table", "copy_worksheet",
                   "delete_worksheet", "rename_worksheet", "merge_cells", "unmerge_cells", "copy_range",
                   "delete_range", "insert_rows", "insert_columns", "delete_sheet_rows", "delete_sheet_columns"}

    @wrap_tool_call
    async def track_excel_writes(request, handler):
        call = request.tool_call
        name, args = call["name"], dict(call.get("args") or {})
        filepath = args.get("filepath")
        if filepath:
            resolved = os.path.realpath(filepath)
            root = os.path.realpath(UPLOAD_DIR)
            if not os.path.isabs(filepath) or os.path.normcase(os.path.dirname(resolved)) != os.path.normcase(root):
                message = f"Error: filepath must be an absolute path directly under {UPLOAD_DIR}"
                errors.append(message)
                return ToolMessage(content=message, tool_call_id=call["id"], status="error")
        async def execute():
            workbook_backup = None
            effective_args = args
            existed_before = bool(filepath and os.path.isfile(filepath))
            try:
                if name in {"write_data_to_excel", "apply_formula", "copy_range", "create_table"} and args.get("sheet_name") in {CHART_SHEET, DATA_SHEET}:
                    raise ValueError(f"{CHART_SHEET} is reserved for charts; use the original source sheet")
                if name == "create_workbook" and os.path.exists(filepath):
                    raise ValueError(f"Workbook already exists at {filepath}; use a new name or edit it")
                if name == "create_pivot_table":
                    if pivot_source and args.get("sheet_name", "").casefold() != pivot_source.casefold():
                        raise ValueError(f"Pivot source must be {pivot_source}, as requested")
                    if re.search(r"\b(total|sum)\b", question, re.I) and args.get("agg_func", "sum").lower() != "sum":
                        raise ValueError("Total pivot values must use SUM")
                chart_count = None
                if name in write_tools and existed_before and filepath.lower().endswith(".xlsx"):
                    fd, workbook_backup = tempfile.mkstemp(suffix=".backup", dir=os.path.dirname(filepath))
                    os.close(fd)
                    shutil.copy2(filepath, workbook_backup)
                if name == "create_chart" and filepath and os.path.isfile(filepath):
                    await asyncio.to_thread(recalculate_formula_caches, filepath)
                    effective_args = prepare_chart_call(filepath, args, question)
                    before = load_workbook(filepath)
                    try:
                        chart_count = len(before[effective_args["sheet_name"]]._charts)
                    finally:
                        before.close()
                effective_request = request.override(tool_call={**call, "args": effective_args}) if effective_args is not args else request
                result = await handler(effective_request)
                content = getattr(result, "content", "")
                if isinstance(content, list):
                    content = " ".join(str(item.get("text", "")) if isinstance(item, dict) else str(item) for item in content)
                if getattr(result, "status", "success") == "error" or str(content).lstrip().lower().startswith("error:"):
                    raise ValueError(str(content))
                if name == "create_pivot_table" and pivot_destination:
                    place_pivot_on_sheet(filepath, args["sheet_name"], pivot_destination, original_sheets)
                    result.content = f"Pivot table saved on {pivot_destination}. Read that sheet for its summary values."
                if name == "create_chart":
                    saved = load_workbook(filepath)
                    try:
                        if chart_count is None or len(saved[effective_args["sheet_name"]]._charts) <= chart_count:
                            raise ValueError("MCP did not save a new chart")
                    finally:
                        saved.close()
                    await asyncio.to_thread(recalculate_formula_caches, filepath)
                    review = verify_chart(filepath, effective_args["sheet_name"], effective_args["data_range"],
                                          effective_args.get("title", ""), effective_args.get("chart_type", ""),
                                          question, args.get("sheet_name", ""))
                    result.content = f"{content}\nSaved chart verification: {review}"
                if name in write_tools and filepath and filepath.lower().endswith(".xlsx"):
                    if not os.path.isfile(filepath):
                        raise ValueError("MCP did not save the workbook")
                    await asyncio.to_thread(recalculate_formula_caches, filepath)
                if name in write_tools:
                    if name == "create_chart":
                        recorded_args = {**effective_args, "sheet_name": CHART_SHEET}
                    elif name == "create_pivot_table" and pivot_destination:
                        recorded_args = {**args, "destination_sheet": pivot_destination}
                    else:
                        recorded_args = effective_args
                    operations.append({"name": name, "args": recorded_args})
                errors[:] = [error for error in errors if not error.startswith(f"{name}:")]
                return result
            except Exception as exc:
                if workbook_backup and os.path.exists(workbook_backup):
                    os.replace(workbook_backup, filepath)
                elif name == "create_workbook" and not existed_before and filepath and os.path.isfile(filepath):
                    os.remove(filepath)
                message = f"{name}: {exc}"
                errors.append(message)
                return ToolMessage(content=message, tool_call_id=call["id"], status="error")
            finally:
                if workbook_backup and os.path.exists(workbook_backup):
                    os.remove(workbook_backup)

        if filepath:
            lock = _workbook_locks.setdefault(os.path.realpath(filepath), asyncio.Lock())
            async with lock:
                return await execute()
        return await execute()

    # 2. Configure OpenRouter model
    model_name = os.getenv("OPENROUTER_MODEL", "nvidia/nemotron-3-super-120b-a12b:free")
    
    llm = ChatOpenRouter(
        model=model_name,
        temperature=0.2,
    )

    import time

    # 3. System prompt: Guide the agent on autonomous inspection & execution
    system_prompt = (
        "You are an expert, proactive, and autonomous Excel AI analyst. You have access to official Excel MCP tools to read, "
        "analyze, inspect, format, and modify Excel workbooks.\n"
        f"Active workbook: '{filename}'. Every filepath passed to a tool must be an absolute path directly under '{UPLOAD_DIR}'. "
        + (f"Use '{file_path}' for the active workbook.\n" if file_path else "If the user asks to create a workbook, call create_workbook with a new .xlsx path there. Otherwise ask for an upload.\n")
    )
    if sheet_name:
        system_prompt += f"Active sheet requested by user: '{sheet_name}'.\n"
    
    system_prompt += (
        "\nCRITICAL OPERATIONAL RULES:\n"
        "1. PROACTIVE & FULL COLUMN INSPECTION:\n"
        "   - NEVER ask the user for column names, letters, rows, or cell coordinates.\n"
        "   - Always discover the sheet layout yourself using 'get_workbook_metadata(include_ranges=True)' to find exact used ranges (e.g. Orders is A1:O501).\n"
        "   - When inspecting headers, ALWAYS read the full column range (e.g. start_cell='A1', end_cell='O1' or 'Z1'). Never declare a column missing without reading the full header span.\n"
        "\n"
        "2. CHART CREATION & PLACEMENT (NEVER OVERLAP DATA):\n"
        "   - NEVER place a chart on top of or inside existing data cells!\n"
        f"   - Every chart is displayed on the dedicated {CHART_SHEET} worksheet. That worksheet must contain chart objects only: never write tables, summaries, metadata, or other cells there.\n"
        "   - Discover the ORIGINAL source sheet, its headers, and full data range. Pass that source sheet and range directly to create_chart. The backend prepares any aggregate summary on a hidden helper sheet, moves the saved chart to Charts, and verifies it.\n"
        "   - Do not create a Charts sheet or prepare a chart summary yourself. For a grouped request, name the chart 'Total <measure> by <category>' so the backend can aggregate all source rows.\n"
        "   - Use chart_type='bar' for a bar request and a valid target_cell; the backend will set the final placement.\n"
        "   - Report success only if the tool response includes Saved chart verification with verified=True. A tool error is a failure.\n"
        "   - For a pivot request, call create_pivot_table on the ORIGINAL source sheet named by the user (for example Orders), with rows and values matching the request and agg_func='sum' for totals. Do not copy raw source rows into Dashboard. The upstream tool creates a temporary <source>_pivot sheet; the backend moves it to a requested new sheet such as Dashboard. Do not create that destination sheet yourself or rename the pivot output. Then call create_chart with the original source sheet and range; the backend prepares the chart data.\n"
        "\n"
        "3. COLOR FORMATTING & aRGB HEX VALUES:\n"
        "   - The 'format_range' tool strictly requires 8-digit uppercase aRGB hex strings (AARRGGBB format, without any '#' prefix).\n"
        "   - NEVER pass text color names like 'red', 'green', 'blue', 'white', or 'black' (they cause tool errors).\n"
        "   - Standard aRGB Hex Codes:\n"
        "     * Red: 'FFFF0000'\n"
        "     * Green: 'FF00B050' or 'FF008000'\n"
        "     * Blue: 'FF0070C0'\n"
        "     * Dark Blue / Navy: 'FF1F4E79' or 'FF002060'\n"
        "     * White: 'FFFFFFFF'\n"
        "     * Black: 'FF000000'\n"
        "     * Yellow: 'FFFFC000' or 'FFFFFF00'\n"
        "     * Gray: 'FF808080' or 'FFD9D9D9'\n"
        "   - Example: To format in bold red, pass bold=True, font_color='FFFF0000'.\n"
        "\n"
        "4. COLUMN INSERTION & SEQUENCING:\n"
        "   - When requested to insert a new column (e.g. 'insert a new column at column G'):\n"
        "     * Call 'insert_columns' FIRST before writing any data or headers into that column!\n"
        "     * Only AFTER inserting, write the header to G1 and data/formulas to G2:G[N].\n"
        "     * If you write data before inserting, the insertion will shift your new data into column H!\n"
        "\n"
        "5. CROSS-SHEET FORMULAS & PREVENTING #VALUE! ERRORS:\n"
        "   - Before writing formulas, inspect what columns actually exist in the target sheet.\n"
        "   - Example in 'Returns' sheet: The columns are OrderID (A), Product (B), ReturnDate (C), Reason (D), RefundAmount (E), Status (F). Notice 'OrderDate' is NOT in 'Returns' — it is in 'Orders' (B)!\n"
        "   - Subtracting text (Product in Column B) from ReturnDate (Column C) produces #VALUE! errors! Never do '=C2-B2'!\n"
        "   - To calculate days between OrderDate (from Orders) and ReturnDate (from Returns):\n"
        "     * Use a cross-sheet lookup formula: '=INT(C2 - VLOOKUP(A2, Orders!$A:$B, 2, FALSE))' or '=INT(C2) - INT(VLOOKUP(A2, Orders!$A:$B, 2, FALSE))'\n"
        "     * OR compute the day differences directly and write the values or formulas using 'write_data_to_excel' with start_cell='G2' for all rows.\n"
        "   - Always populate all rows corresponding to the data (e.g., rows 2 to 36 for Returns).\n"
        "\n"
        "6. BATCH WRITING VS SINGLE CELL:\n"
        "   - Use 'apply_formula' when applying a formula to a single cell (e.g. P3).\n"
        "   - Use 'write_data_to_excel' when writing multiple rows of values or formulas at once (pass a list of rows, e.g. [[val1], [val2], ...] with start_cell='G2').\n"
        "\n"
        "7. AUTONOMOUS MULTI-TOOL EXECUTION:\n"
        "   - Execute all required steps in order (inspect -> insert -> write -> formula -> format) without stopping halfway.\n"
        "   - After completing all operations, answer in one short sentence describing the result. For a column edit, name the changed column. Do not list steps, tool names, raw paths, arguments, metadata, or verification details.\n"
    )

    # 4. Create LangChain agent with MCP tools, checkpointer, and memory store
    agent = create_agent(
        model=llm,
        tools=mcp_tools,
        system_prompt=system_prompt,
        checkpointer=_checkpointer,
        store=_store,
        middleware=[track_excel_writes, retry_provider_errors],
    )

    from langchain_core.messages import HumanMessage
    
    print(f"\n[Agent] Starting query execution for thread: {session_id}")
    start_time = time.time()
    
    # 5. Invoke the agent with thread memory
    try:
        response = await agent.ainvoke(
            {
                "messages": [
                    HumanMessage(content=question)
                ]
            },
            config={"configurable": {"thread_id": session_id}}
        )
    except Exception as e:
        elapsed = time.time() - start_time
        if "provider_overloaded" in str(e).lower():
            print(f"[Agent] Model provider overloaded after retries ({elapsed:.2f}s)")
            message = "The model provider is temporarily overloaded. Please try again shortly."
            return {
                "answer": message,
                "chart_image": None,
                "operations": operations,
                "tool_errors": [*errors, message],
            }
        print(f"[Agent] Error during execution ({elapsed:.2f}s): {e}")
        if "provider_unavailable" in str(e):
            message = "OpenRouter's Nemotron 3 Super free provider is temporarily unavailable. Please try again shortly."
            return {
                "answer": message,
                "chart_image": None,
                "operations": operations,
                "tool_errors": [*errors, message],
            }
        # Detect Quota error specifically
        if "429" in str(e) or "Quota" in str(e):
            return {
                "answer": "⚠️ **API Quota Exceeded:** I hit my rate limit for the free tier while processing your request. Please wait a minute and try again, or switch to a higher tier API key.",
                "chart_image": None,
                "operations": operations,
                "tool_errors": errors,
            }
        raise e

    elapsed = time.time() - start_time
    print(f"[Agent] Execution completed in {elapsed:.2f} seconds.")

    # 6. Extract the final response text
    messages = response.get("messages", [])
    if messages:
        last_message = messages[-1]
        content = getattr(last_message, "content", str(last_message))
        if isinstance(content, list):
            parts = []
            for item in content:
                if isinstance(item, str):
                    parts.append(item)
                elif isinstance(item, dict) and "text" in item:
                    parts.append(item["text"])
                else:
                    parts.append(str(item))
            answer = "\n".join(parts)
        elif isinstance(content, dict):
            answer = content.get("text", str(content))
        else:
            answer = str(content)
    else:
        answer = "No response generated."

    print(f"[Agent] Final Response: {answer[:200]}...\n")

    return {
        "answer": answer,
        "chart_image": None,
        "operations": operations,
        "tool_errors": errors,
    }


def ask_agent(
    question: str,
    file_path: Optional[str] = None,
    sheet_name: Optional[str] = None,
    thread_id: Optional[str] = "default_session",
) -> dict:
    """Synchronous wrapper for aask_agent."""
    try:
        loop = asyncio.get_running_loop()
    except RuntimeError:
        loop = None

    if loop and loop.is_running():
        import concurrent.futures
        with concurrent.futures.ThreadPoolExecutor() as pool:
            return pool.submit(asyncio.run, aask_agent(question, file_path, sheet_name, thread_id)).result()
    else:
        return asyncio.run(aask_agent(question, file_path, sheet_name, thread_id))
