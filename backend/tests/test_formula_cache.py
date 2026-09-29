import os
import tempfile
import unittest
import zipfile
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

from agents.chart_prepare import prepare_chart_call
from agents.chart_review import verify_chart
from agents.formula_cache import recalculate_formula_caches
from workbook_sync import updated_tables


class FormulaCacheTests(unittest.TestCase):
    def test_arithmetic_chain_and_existing_chart_survive_recalculation(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "book.xlsx")
            book = Workbook()
            sheet = book.active
            sheet.title = "Orders"
            sheet.append(["Quantity", "Annual_Target", "Annual_Target_x2"])
            sheet.append([21, "=A2*2", "=B2*2"])
            chart = BarChart()
            chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2), titles_from_data=True)
            sheet.add_chart(chart, "E2")
            book.save(path)
            book.close()

            with zipfile.ZipFile(path) as archive:
                drawings = {name: archive.read(name) for name in archive.namelist()
                            if name.startswith("xl/charts/") or name.startswith("xl/drawings/")}
            self.assertTrue(recalculate_formula_caches(path))
            with zipfile.ZipFile(path) as archive:
                self.assertEqual(drawings, {name: archive.read(name) for name in drawings})
            values = load_workbook(path, data_only=True)
            formulas = load_workbook(path, data_only=False)
            try:
                self.assertEqual([values["Orders"]["B2"].value, values["Orders"]["C2"].value], [42, 84])
                self.assertEqual(formulas["Orders"]["C2"].value, "=B2*2")
                self.assertEqual(len(formulas["Orders"]._charts), 1)
                preview = updated_tables(directory, [{"name": "apply_formula", "args": {
                    "filepath": path, "sheet_name": "Orders", "cell": "C2"}}])
                self.assertEqual(preview[0]["rows"][0], ["84"])
            finally:
                values.close()
                formulas.close()

    def test_cross_sheet_vlookup_with_whole_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "book.xlsx")
            book = Workbook()
            orders = book.active
            orders.title = "Orders"
            orders.append(["ID", "Price"])
            orders.append(["A", 21])
            returns = book.create_sheet("Returns")
            returns.append(["ID", "Price", "Double"])
            returns.append(["A", '=VLOOKUP(A2,orders!A:B,2,FALSE)', "=B2*2"])
            book.save(path)
            book.close()

            self.assertTrue(recalculate_formula_caches(path))
            cached = load_workbook(path, data_only=True)
            formulas = load_workbook(path, data_only=False)
            try:
                self.assertEqual(cached["Returns"]["B2"].value, 21)
                self.assertEqual(cached["Returns"]["C2"].value, 42)
                self.assertEqual(formulas["Returns"]["B2"].value, '=VLOOKUP(A2,orders!A:B,2,FALSE)')
            finally:
                cached.close()
                formulas.close()

    def test_formula_error_does_not_replace_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "book.xlsx")
            book = Workbook()
            book.active["A1"] = "=1/0"
            book.save(path)
            book.close()
            with open(path, "rb") as stream:
                original = stream.read()
            with self.assertRaises(ValueError):
                recalculate_formula_caches(path)
            with open(path, "rb") as stream:
                self.assertEqual(stream.read(), original)

    def test_date_formula_keeps_excel_date_format(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "book.xlsx")
            book = Workbook()
            book.active["A1"] = "=DATE(2024,1,1)"
            book.active["A1"].number_format = "yyyy-mm-dd"
            book.save(path)
            book.close()
            recalculate_formula_caches(path)
            cached = load_workbook(path, data_only=True)
            self.assertEqual(cached.active["A1"].value, datetime(2024, 1, 1))
            cached.close()

    def test_grouped_chart_checks_formula_backed_source_values(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "book.xlsx")
            book = Workbook()
            orders = book.active
            orders.title = "Orders"
            orders.append(["Region", "Base", "Sales"])
            orders.append(["Europe", 10, "=B2*2"])
            orders.append(["Asia", 5, "=B3*2"])
            orders.append(["Europe", 3, "=B4*2"])
            book.save(path)
            book.close()
            recalculate_formula_caches(path)
            args = prepare_chart_call(path, {
                "sheet_name": "Orders", "data_range": "A1:C4", "title": "Total Sales by Region",
                "chart_type": "bar", "target_cell": "G2"},
                "Create a bar chart showing total Sales by Region from Orders")
            book = load_workbook(path)
            charts = book["_ChartData"]
            self.assertEqual(charts["B2"].value, 26)
            chart = BarChart()
            chart.title = args["title"]
            chart.add_data(Reference(charts, min_col=2, min_row=1, max_row=3), titles_from_data=True)
            chart.set_categories(Reference(charts, min_col=1, min_row=2, max_row=3))
            charts.add_chart(chart, args["target_cell"])
            book.save(path)
            book.close()
            recalculate_formula_caches(path)
            result = verify_chart(path, "_ChartData", args["data_range"], args["title"], "bar",
                                  "Create a bar chart showing total Sales by Region from Orders")
            self.assertTrue(result["verified"])
            recalculate_formula_caches(path)
            book = load_workbook(path, data_only=True)
            self.assertEqual(book["Orders"]["C2"].value, 20)
            book.close()


if __name__ == "__main__":
    unittest.main()
