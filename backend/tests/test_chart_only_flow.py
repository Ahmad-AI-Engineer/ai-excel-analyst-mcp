import os
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

from agents.chart_prepare import prepare_chart_call
from agents.chart_review import CHART_SHEET, DATA_SHEET, _formula, _rect, verify_chart


def save_mcp_chart(path, args):
    book = load_workbook(path)
    sheet = book[args["sheet_name"]]
    top = int(args["data_range"].split(":")[0][1:])
    bottom = int(args["data_range"].split(":")[1][1:])
    chart = BarChart()
    chart.title = args["title"]
    chart.add_data(Reference(sheet, min_col=2, min_row=top, max_row=bottom), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=top + 1, max_row=bottom))
    sheet.add_chart(chart, args["target_cell"])
    book.save(path)
    book.close()


class ChartOnlyFlowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.temporary.name, "book.xlsx")
        book = Workbook()
        source = book.active
        source.title = "Orders"
        source.append(["Region", "Sales"])
        source.append(["Europe", 10])
        source.append(["Asia", 20])
        source.append(["Europe", 5])
        book.save(self.path)
        book.close()

    def tearDown(self):
        self.temporary.cleanup()

    def create(self, title="Total Sales by Region", question="Create a bar chart showing total Sales by Region from Orders"):
        args = prepare_chart_call(self.path, {
            "sheet_name": "Orders", "data_range": "A1:B4", "chart_type": "bar",
            "title": title, "target_cell": "G2"}, question)
        self.assertEqual(args["sheet_name"], DATA_SHEET)
        save_mcp_chart(self.path, args)
        return verify_chart(self.path, DATA_SHEET, args["data_range"], title, "bar", question, "Orders")

    def test_chart_sheet_has_only_chart_objects(self):
        result = self.create()
        self.assertEqual(result["sheet"], CHART_SHEET)
        book = load_workbook(self.path)
        try:
            self.assertEqual(len(book[CHART_SHEET]._charts), 1)
            self.assertFalse(any(cell.value is not None for cell in book[CHART_SHEET]._cells.values()))
            self.assertIs(book[CHART_SHEET].sheet_view.showGridLines, False)
            self.assertEqual(book[DATA_SHEET].sheet_state, "hidden")
            chart = book[CHART_SHEET]._charts[0]
            self.assertIs(chart.visible_cells_only, False)
            self.assertIn(DATA_SHEET, _formula(chart.series[0].cat))
            self.assertIn(DATA_SHEET, _formula(chart.series[0].val))
            self.assertEqual(book["Orders"]["B4"].value, 5)
        finally:
            book.close()

    def test_second_chart_and_update_do_not_overlap_or_duplicate(self):
        first = self.create()
        second = self.create("Sales by Region", "Create another bar chart of Sales by Region from Orders")
        self.assertGreaterEqual(second["bounds"][1], first["bounds"][3])
        book = load_workbook(self.path)
        self.assertEqual(len(book[CHART_SHEET]._charts), 2)
        helper_rows = book[DATA_SHEET].max_row
        book.close()
        self.create()
        book = load_workbook(self.path)
        try:
            self.assertEqual(len(book[CHART_SHEET]._charts), 2)
            self.assertFalse(any(cell.value is not None for cell in book[CHART_SHEET]._cells.values()))
            self.assertEqual(book[DATA_SHEET].max_row, helper_rows)
        finally:
            book.close()

    def test_existing_chart_cells_are_archived_without_data_loss(self):
        book = load_workbook(self.path)
        old = book.create_sheet(CHART_SHEET)
        old.append(["Title", "Old Sales"])
        old.append(["Source", "Orders!A1:B4"])
        old.append(["Request", "SUM(Sales) by Region"])
        old.append(["Type", "bar"])
        old["D1"], old["E1"] = "Region", "Sum of Sales"
        old["D2"], old["E2"] = "Europe", 15
        old["D3"], old["E3"] = "Asia", 20
        old_chart = BarChart()
        old_chart.title = "Old Sales"
        old_chart.add_data(Reference(old, min_col=5, min_row=1, max_row=3), titles_from_data=True)
        old_chart.set_categories(Reference(old, min_col=4, min_row=2, max_row=3))
        old.add_chart(old_chart, "G2")
        book.save(self.path)
        book.close()

        self.create()
        book = load_workbook(self.path)
        try:
            self.assertEqual(book["_ChartArchive"].sheet_state, "hidden")
            self.assertEqual(book["_ChartArchive"]["B2"].value, "Orders!A1:B4")
            self.assertEqual(book["_ChartArchive"]["E3"].value, 20)
            self.assertEqual(len(book[CHART_SHEET]._charts), 2)
            self.assertFalse(any(cell.value is not None for cell in book[CHART_SHEET]._cells.values()))
            old_chart = next(chart for chart in book[CHART_SHEET]._charts if _rect(chart)[1] == 2)
            self.assertIn("_ChartArchive", _formula(old_chart.series[0].cat))
            self.assertIs(old_chart.visible_cells_only, False)
        finally:
            book.close()

    def test_matching_legacy_chart_is_replaced_even_when_source_quoting_differs(self):
        book = load_workbook(self.path)
        old = book.create_sheet(CHART_SHEET)
        for row in (["Region", "Sum of Sales", None, "Title", "Total Sales by Region"],
                    ["Europe", 15, None, "Source", "Orders!A1:B4"],
                    ["Asia", 20, None, "Request", "SUM(Sales) by Region"],
                    [None, None, None, "Type", "bar"]):
            old.append(row)
        chart = BarChart()
        chart.title = "Total Sales by Region"
        chart.add_data(Reference(old, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        chart.set_categories(Reference(old, min_col=1, min_row=2, max_row=3))
        old.add_chart(chart, "G2")
        book.save(self.path)
        book.close()

        self.create()
        book = load_workbook(self.path)
        try:
            self.assertEqual(len(book[CHART_SHEET]._charts), 1)
            self.assertEqual(book["_ChartArchive"]["E2"].value, "Orders!A1:B4")
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
