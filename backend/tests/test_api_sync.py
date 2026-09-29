import io
import os
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from openpyxl import Workbook

import main
from agents.formula_cache import recalculate_formula_caches


def workbook_bytes():
    stream = io.BytesIO()
    book = Workbook()
    book.active.title = "Orders"
    book.active.append(["Region", "Sales"])
    book.active.append(["Europe", 10])
    book.save(stream)
    book.close()
    return stream.getvalue()


class ApiSyncTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = self.temp.name
        self.patch_dir = patch.object(main, "UPLOAD_DIR", self.directory)
        self.patch_dir.start()
        self.client = TestClient(main.app)

    def tearDown(self):
        self.patch_dir.stop()
        self.temp.cleanup()

    def test_create_without_upload_and_preview(self):
        path = os.path.join(self.directory, "new.xlsx")

        async def agent(**kwargs):
            self.assertIsNone(kwargs["file_path"])
            with open(path, "wb") as stream:
                stream.write(workbook_bytes())
            return {"answer": "Created workbook", "chart_image": None,
                    "operations": [{"name": "create_workbook", "args": {"filepath": path}}], "tool_errors": []}

        with patch.object(main, "aask_agent", side_effect=agent) as call:
            response = self.client.post("/chat", json={"question": "Create a workbook"})
        self.assertEqual(call.call_count, 1)
        data = response.json()
        self.assertEqual(data["created_files"], ["new.xlsx"])
        self.assertEqual(self.client.get("/files").json()["files"], ["new.xlsx"])
        preview = self.client.get("/preview/new.xlsx?sheet_name=Orders&cell_range=A1:B2").json()
        self.assertEqual(preview["columns"], ["Region", "Sales"])
        self.assertEqual(preview["data"][0]["Sales"], 10)

    def test_formula_results_show_in_full_and_ranged_previews(self):
        path = os.path.join(self.directory, "formulas.xlsx")
        book = Workbook()
        book.active.title = "Orders"
        book.active.append(["Amount", "Double"])
        book.active.append([21, "=A2*2"])
        book.save(path)
        book.close()
        recalculate_formula_caches(path)
        for suffix in ("", "?sheet_name=Orders&cell_range=A1:B2"):
            data = self.client.get("/preview/formulas.xlsx" + suffix).json()
            self.assertEqual(data["data"][0]["Double"], 42)

    def test_edit_uploaded_table_and_handle_preview_failure_once(self):
        self.client.post("/upload", files={"file": ("uploaded.xlsx", workbook_bytes())})
        path = os.path.join(self.directory, "uploaded.xlsx")

        async def agent(**kwargs):
            self.assertEqual(kwargs["file_path"], path)
            from openpyxl import load_workbook
            book = load_workbook(path)
            book["Orders"]["B2"] = 12
            book.save(path)
            book.close()
            return {"answer": "Updated", "chart_image": None,
                    "operations": [{"name": "write_data_to_excel", "args": {
                        "filepath": path, "sheet_name": "Orders", "start_cell": "B2", "data": [[12]]}}],
                    "tool_errors": []}

        with patch.object(main, "aask_agent", side_effect=agent) as call:
            response = self.client.post("/chat", json={"question": "Change the sales", "file_name": "uploaded.xlsx"})
        data = response.json()
        self.assertEqual(call.call_count, 1)
        self.assertEqual(data["updated_tables"][0]["rows"][0], ["12"])
        self.assertEqual(data["created_files"], [])

        with patch.object(main, "aask_agent", side_effect=agent) as call, patch.object(
                main, "updated_tables", side_effect=RuntimeError("preview offline")):
            response = self.client.post("/chat", json={"question": "Change the sales", "file_name": "uploaded.xlsx"})
        self.assertEqual(call.call_count, 1)
        self.assertIn("preview offline", response.json()["sync_error"])

    def test_new_workbook_with_existing_upload_keeps_original(self):
        self.client.post("/upload", files={"file": ("old.xlsx", workbook_bytes())})
        original_path = os.path.join(self.directory, "old.xlsx")
        new_path = os.path.join(self.directory, "new.xlsx")

        async def agent(**kwargs):
            self.assertEqual(kwargs["file_path"], original_path)
            with open(new_path, "wb") as stream:
                stream.write(workbook_bytes())
            return {"answer": "Created another workbook", "chart_image": None,
                    "operations": [{"name": "create_workbook", "args": {"filepath": new_path}}],
                    "tool_errors": []}

        with patch.object(main, "aask_agent", side_effect=agent):
            response = self.client.post("/chat", json={"question": "Create another workbook", "file_name": "old.xlsx"})
        self.assertEqual(response.json()["created_files"], ["new.xlsx"])
        self.assertEqual(self.client.get("/files").json()["files"], ["old.xlsx", "new.xlsx"])
        self.assertEqual(self.client.get("/preview/old.xlsx").json()["data"][0]["Sales"], 10)

    def test_short_write_reply_and_failed_chart_do_not_claim_success(self):
        self.client.post("/upload", files={"file": ("uploaded.xlsx", workbook_bytes())})
        path = os.path.join(self.directory, "uploaded.xlsx")

        async def edit_agent(**kwargs):
            return {"answer": "**Steps performed**\n| Tool | Details |\nA long transcript", "operations": [
                {"name": "write_data_to_excel", "args": {"filepath": path, "sheet_name": "Orders",
                 "start_cell": "B2", "data": [[12]]}}], "tool_errors": []}

        with patch.object(main, "aask_agent", side_effect=edit_agent):
            data = self.client.post("/chat", json={"question": "Change Sales", "file_name": "uploaded.xlsx"}).json()
        self.assertEqual(data["answer"], "Updated Orders: Sales.")
        self.assertEqual(data["updated_tables"][0]["headers"], ["Sales"])

        async def chart_agent(**kwargs):
            return {"answer": "Chart created and verified", "operations": [
                {"name": "write_data_to_excel", "args": {"filepath": path, "sheet_name": "Orders",
                 "start_cell": "D1", "data": [["Title", "Bad chart"], ["Source", "Orders!A1:B2"],
                                              ["Request", "SUM(Sales) by Region"], ["Type", "pie"]]}}],
                    "tool_errors": ["create_chart: Saved chart series mismatch"]}

        with patch.object(main, "aask_agent", side_effect=chart_agent):
            data = self.client.post("/chat", json={"question": "Create chart", "file_name": "uploaded.xlsx"}).json()
        self.assertIn("couldn't finish", data["answer"])
        self.assertNotIn("created and verified", data["answer"])
        self.assertEqual(data["updated_tables"], [])

    def test_partial_pivot_does_not_claim_chart_success(self):
        self.client.post("/upload", files={"file": ("uploaded.xlsx", workbook_bytes())})

        async def agent(**kwargs):
            return {"answer": "Created the pivot and pie chart", "operations": [
                {"name": "create_pivot_table", "args": {"sheet_name": "Orders", "destination_sheet": "Dashboard"}}],
                    "tool_errors": []}

        with patch.object(main, "aask_agent", side_effect=agent):
            data = self.client.post("/chat", json={"question": "Create a pivot and add a pie chart",
                                                   "file_name": "uploaded.xlsx"}).json()
        self.assertIn("No verified chart was saved", data["answer"])


if __name__ == "__main__":
    unittest.main()
