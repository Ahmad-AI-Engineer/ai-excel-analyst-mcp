import os
import tempfile
import unittest
from unittest.mock import patch

from langchain_core.messages import AIMessage, ToolMessage
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table

from agents import graph


class FakeClient:
    async def get_tools(self):
        return []


class FakeRequest:
    def __init__(self, tool_call):
        self.tool_call = tool_call

    def override(self, **changes):
        return FakeRequest(changes.get("tool_call", self.tool_call))


class ToolTrackingTests(unittest.IsolatedAsyncioTestCase):
    async def run_call(self, path, name, args, handler, question="test"):
        captured = {}

        def make_agent(**kwargs):
            middleware = kwargs["middleware"][0]

            class Agent:
                async def ainvoke(self, *_args, **_kwargs):
                    request = FakeRequest({"id": "test-id", "name": name, "args": args})
                    captured["message"] = await middleware.awrap_tool_call(request, handler)
                    return {"messages": [AIMessage(content="done")]}

            return Agent()

        with patch.object(graph, "UPLOAD_DIR", os.path.dirname(path)), patch.object(
                graph, "get_mcp_client", return_value=FakeClient()), patch.object(
                graph, "ChatOpenRouter", return_value=object()), patch.object(
                graph, "create_agent", side_effect=make_agent):
            result = await graph.aask_agent(question, file_path=path)
        return result, captured["message"]

    async def test_failed_chart_save_is_not_recorded(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "book.xlsx")
            book = Workbook()
            book.active.append(["Region", "Sales"])
            book.active.append(["Europe", 10])
            book.active.append(["Asia", 20])
            book.save(path)
            book.close()

            async def handler(request):
                return ToolMessage(content="Bar chart created successfully", tool_call_id="test-id")

            result, message = await self.run_call(path, "create_chart", {
                "filepath": path, "sheet_name": "Sheet", "data_range": "A1:B3",
                "title": "Sales", "chart_type": "bar", "target_cell": "D2"}, handler)
            self.assertEqual(message.status, "error")
            self.assertIn("did not save", message.content)
            self.assertEqual(result["operations"], [])

    async def test_existing_workbook_is_not_recreated(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "book.xlsx")
            Workbook().save(path)
            called = False

            async def handler(request):
                nonlocal called
                called = True
                return ToolMessage(content="Created", tool_call_id="test-id")

            result, message = await self.run_call(path, "create_workbook", {"filepath": path}, handler)
            self.assertFalse(called)
            self.assertEqual(message.status, "error")
            self.assertEqual(result["operations"], [])

    async def test_agent_cannot_write_details_on_chart_sheet(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "book.xlsx")
            Workbook().save(path)
            called = False

            async def handler(request):
                nonlocal called
                called = True
                return ToolMessage(content="Wrote data", tool_call_id="test-id")

            result, message = await self.run_call(path, "write_data_to_excel", {
                "filepath": path, "sheet_name": "Charts", "start_cell": "A1",
                "data": [["Title", "A chart"]]}, handler)
            self.assertFalse(called)
            self.assertEqual(message.status, "error")
            self.assertEqual(result["operations"], [])

    async def test_successful_write_records_arguments(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "book.xlsx")
            Workbook().save(path)
            args = {"filepath": path, "sheet_name": "Sheet", "start_cell": "A1", "data": [["Region"]]}

            async def handler(request):
                return ToolMessage(content="Wrote data", tool_call_id="test-id")

            result, message = await self.run_call(path, "write_data_to_excel", args, handler)
            self.assertEqual(message.status, "success")
            self.assertEqual(result["operations"], [{"name": "write_data_to_excel", "args": args}])

    async def test_formula_write_is_cached_and_failed_formula_is_rolled_back(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "book.xlsx")
            book = Workbook()
            book.active.append(["Amount", "Double"])
            book.active.append([21, None])
            book.save(path)
            book.close()
            args = {"filepath": path, "sheet_name": "Sheet", "start_cell": "B2", "data": [["=A2*2"]]}

            async def write_formula(request):
                book = load_workbook(path)
                book.active["B2"] = request.tool_call["args"]["data"][0][0]
                book.save(path)
                book.close()
                return ToolMessage(content="Wrote data", tool_call_id="test-id")

            result, message = await self.run_call(path, "write_data_to_excel", args, write_formula)
            self.assertEqual(message.status, "success", message.content)
            self.assertEqual(len(result["operations"]), 1)
            cached = load_workbook(path, data_only=True)
            self.assertEqual(cached.active["B2"].value, 42)
            cached.close()
            with open(path, "rb") as stream:
                previous = stream.read()
            bad_args = {**args, "data": [["=1/0"]]}
            result, message = await self.run_call(path, "write_data_to_excel", bad_args, write_formula)
            self.assertEqual(message.status, "error")
            self.assertEqual(result["operations"], [])
            with open(path, "rb") as stream:
                self.assertEqual(stream.read(), previous)

    async def test_failed_chart_verification_restores_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "book.xlsx")
            book = Workbook()
            sheet = book.active
            sheet.append(["Region", "Sales", "Wrong category"])
            sheet.append(["Europe", 10, "A"])
            sheet.append(["Asia", 20, "B"])
            book.save(path)
            book.close()
            with open(path, "rb") as stream:
                before = stream.read()

            async def handler(request):
                book = load_workbook(path)
                sheet = book[request.tool_call["args"]["sheet_name"]]
                chart = BarChart()
                chart.title = "Sales"
                chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
                chart.set_categories(Reference(sheet, min_col=4, min_row=2, max_row=3))
                sheet.add_chart(chart, "A2")
                book.save(path)
                book.close()
                return ToolMessage(content="Created", tool_call_id="test-id")

            result, message = await self.run_call(path, "create_chart", {
                "filepath": path, "sheet_name": "Sheet", "data_range": "A1:B3",
                "title": "Sales", "chart_type": "bar", "target_cell": "D2"}, handler)
            self.assertEqual(message.status, "error")
            self.assertIn("series mismatch", message.content)
            self.assertEqual(result["operations"], [])
            with open(path, "rb") as stream:
                self.assertEqual(stream.read(), before)

    async def test_explicit_orders_chart_is_routed_to_charts(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "book.xlsx")
            book = Workbook()
            orders = book.active
            orders.title = "Orders"
            orders.append(["Region", "Sales"])
            orders.append(["Europe", 10])
            orders.append(["Asia", 20])
            orders["H1"], orders["I1"] = "Region", "Sum of Sales"
            orders["H2"], orders["I2"] = "Europe", 10
            orders["H3"], orders["I3"] = "Asia", 20
            old_chart = BarChart()
            old_chart.title = "Total Sales by Region"
            old_chart.add_data(Reference(orders, min_col=2, min_row=1, max_row=3), titles_from_data=True)
            old_chart.set_categories(Reference(orders, min_col=1, min_row=2, max_row=3))
            orders.add_chart(old_chart, "D2")
            book.save(path)
            book.close()

            async def handler(request):
                routed = request.tool_call["args"]
                self.assertEqual(routed["sheet_name"], "_ChartData")
                book = load_workbook(path)
                sheet = book["_ChartData"]
                chart = BarChart()
                chart.title = routed["title"]
                chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
                chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
                sheet.add_chart(chart, routed["target_cell"])
                book.save(path)
                book.close()
                return ToolMessage(content="Created", tool_call_id="test-id")

            result, message = await self.run_call(path, "create_chart", {
                "filepath": path, "sheet_name": "Orders", "data_range": "H1:I3",
                "title": "Total Sales by Region", "chart_type": "bar", "target_cell": "K2"},
                handler, "Create a bar chart in Orders showing total Sales by Region")
            self.assertEqual(message.status, "success")
            self.assertEqual(result["operations"][-1]["args"]["sheet_name"], "Charts")
            book = load_workbook(path)
            self.assertEqual(len(book["Orders"]._charts), 0)
            self.assertEqual(len(book["Charts"]._charts), 1)
            self.assertEqual(book["_ChartData"].sheet_state, "hidden")
            self.assertFalse(any(cell.value is not None for cell in book["Charts"]._cells.values()))
            self.assertEqual(book["Orders"]["B3"].value, 20)
            book.close()

    async def test_raw_table_chart_call_aggregates_and_replaces_matching_chart(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "book.xlsx")
            book = Workbook()
            orders = book.active
            orders.title = "Orders"
            for row in (["Region", "Sales", "Segment"], ["Europe", 10, "A"],
                        ["Asia", 20, "B"], ["Europe", 5, "C"]):
                orders.append(row)
            book.save(path)
            book.close()

            async def handler(request):
                routed = request.tool_call["args"]
                self.assertEqual(routed["sheet_name"], "_ChartData")
                book = load_workbook(path)
                sheet = book["_ChartData"]
                self.assertEqual([sheet["A1"].value, sheet["B1"].value], ["Region", "Sum of Sales"])
                self.assertEqual([sheet["A2"].value, sheet["B2"].value], ["Europe", 15])
                self.assertEqual([sheet["A3"].value, sheet["B3"].value], ["Asia", 20])
                self.assertEqual(sheet["E2"].value, "'Orders'!A1:C4")
                start = int(routed["data_range"].split(":")[0][1:])
                finish = int(routed["data_range"].split(":")[1][1:])
                chart = BarChart()
                chart.title = routed["title"]
                chart.add_data(Reference(sheet, min_col=2, min_row=start, max_row=finish), titles_from_data=True)
                chart.set_categories(Reference(sheet, min_col=1, min_row=start + 1, max_row=finish))
                sheet.add_chart(chart, routed["target_cell"])
                book.save(path)
                book.close()
                return ToolMessage(content="Created", tool_call_id="test-id")

            args = {"filepath": path, "sheet_name": "Orders", "data_range": "A1:C4",
                    "title": "Total Sales by Region", "chart_type": "bar", "target_cell": "Q2"}
            question = "Create a bar chart on Orders showing total Sales by Region"
            first, first_message = await self.run_call(path, "create_chart", args, handler, question)
            self.assertEqual(first_message.status, "success")
            second, second_message = await self.run_call(path, "create_chart", args, handler, question)
            self.assertEqual(second_message.status, "success", second_message.content)
            book = load_workbook(path)
            self.assertEqual(len(book["Charts"]._charts), 1)
            self.assertFalse(any(cell.value is not None for cell in book["Charts"]._cells.values()))
            self.assertEqual(len(book["Orders"]._charts), 0)
            self.assertEqual(book["Orders"]["C4"].value, "C")
            book.close()

    async def test_salary_chart_repairs_summary_source_and_uses_charts_sheet(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "book.xlsx")
            book = Workbook()
            employees = book.active
            employees.title = "Employees"
            employees.append(["Department", "Salary"])
            employees.append(["Engineering", 2])
            employees.append(["Sales", 7])
            employees.append(["Engineering", 3])
            summary = book.create_sheet("SalaryChart")
            summary.append(["Department", "Total Salary", None, "Title", "Total Salary by Department"])
            summary.append(["Engineering", 5, None, "Source", "SalaryChart!A1:B3"])
            summary.append(["Sales", 7, None, "Request", "SUM(Salary) by Department"])
            summary["D4"], summary["E4"] = "Type", "bar"
            book.save(path)
            book.close()

            async def handler(request):
                routed = request.tool_call["args"]
                self.assertEqual(routed["sheet_name"], "_ChartData")
                book = load_workbook(path)
                charts = book["_ChartData"]
                chart = BarChart()
                chart.title = routed["title"]
                chart.add_data(Reference(charts, min_col=2, min_row=1, max_row=3), titles_from_data=True)
                chart.set_categories(Reference(charts, min_col=1, min_row=2, max_row=3))
                charts.add_chart(chart, routed["target_cell"])
                book.save(path)
                book.close()
                return ToolMessage(content="Created", tool_call_id="test-id")

            result, message = await self.run_call(path, "create_chart", {
                "filepath": path, "sheet_name": "SalaryChart", "data_range": "A1:B3",
                "title": "Total Salary by Department", "chart_type": "bar", "target_cell": "G2"},
                handler, "Create a bar chart of total Salary by Department from Employees in SalaryChart")
            self.assertEqual(message.status, "success")
            book = load_workbook(path)
            self.assertEqual(book["_ChartData"]["E2"].value, "'Employees'!A1:B4")
            self.assertEqual(len(book["Charts"]._charts), 1)
            self.assertFalse(any(cell.value is not None for cell in book["Charts"]._cells.values()))
            self.assertEqual(len(book["SalaryChart"]._charts), 0)
            book.close()

    async def test_pivot_uses_orders_and_moves_output_to_dashboard(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "book.xlsx")
            book = Workbook()
            book.active.title = "Orders"
            book.active.append(["Category", "Profit"])
            book.active.append(["Hardware", 10])
            book.save(path)
            book.close()
            question = 'Create a new sheet called "Dashboard". Create a pivot table of total Profit by Category from Orders.'

            async def handler(request):
                book = load_workbook(path)
                repeated_call = "Dashboard" in book
                if "Dashboard" not in book:
                    book.create_sheet("Dashboard")  # An empty sheet created earlier by the agent.
                pivot = book.create_sheet("Orders_pivot")
                pivot.append(["Category", "Profit (sum)"])
                pivot.append(["Hardware", 10])
                pivot.add_table(Table(displayName="PivotTable_Test2" if repeated_call else "PivotTable_Test", ref="A1:B2"))
                book.save(path)
                book.close()
                return ToolMessage(content="Pivot created", tool_call_id="test-id")

            result, message = await self.run_call(path, "create_pivot_table", {
                "filepath": path, "sheet_name": "Orders", "data_range": "A1:B2",
                "rows": ["Category"], "values": ["Profit"], "agg_func": "sum"}, handler, question)
            self.assertEqual(message.status, "success", message.content)
            self.assertIn("Dashboard", message.content)
            self.assertEqual(result["operations"][0]["args"]["destination_sheet"], "Dashboard")
            book = load_workbook(path)
            self.assertEqual(book.sheetnames, ["Orders", "Dashboard"])
            self.assertEqual(book["Dashboard"]["B2"].value, 10)
            self.assertEqual(book["Orders"]["B2"].value, 10)
            book.close()
            repeated, repeated_message = await self.run_call(path, "create_pivot_table", {
                "filepath": path, "sheet_name": "Orders", "data_range": "A1:B2",
                "rows": ["Category"], "values": ["Profit"], "agg_func": "sum"}, handler, question)
            self.assertEqual(repeated_message.status, "success")
            book = load_workbook(path)
            self.assertEqual(book.sheetnames, ["Orders", "Dashboard"])
            book.close()


if __name__ == "__main__":
    unittest.main()
