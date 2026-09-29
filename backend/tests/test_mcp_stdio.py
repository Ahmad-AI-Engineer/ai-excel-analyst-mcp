"""Run with RUN_MCP_INTEGRATION=1 in the backend image."""

import os
import tempfile
import unittest

from openpyxl import load_workbook


@unittest.skipUnless(os.getenv("RUN_MCP_INTEGRATION") == "1", "requires pinned stdio server")
class StdioIntegrationTests(unittest.IsolatedAsyncioTestCase):
    async def test_discover_and_write_workbook(self):
        from agents.graph import get_mcp_client

        tools = {tool.name: tool for tool in await get_mcp_client().get_tools()}
        self.assertTrue({"create_workbook", "create_worksheet", "write_data_to_excel", "create_chart"} <= tools.keys())
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "created.xlsx")

            async def call(name, **args):
                result = await tools[name].ainvoke(args)
                self.assertFalse(str(result).lower().startswith("error:"), result)
                return result

            await call("create_workbook", filepath=path)
            await call("create_worksheet", filepath=path, sheet_name="Orders")
            await call("write_data_to_excel", filepath=path, sheet_name="Orders",
                       start_cell="A1", data=[["Region", "Sales"], ["Europe", 4], ["Europe", 6], ["Asia", 8]])
            from agents.chart_prepare import prepare_chart_call
            routed = prepare_chart_call(path, {"sheet_name": "Orders", "data_range": "A1:B4",
                                               "chart_type": "bar", "target_cell": "G2",
                                               "title": "Total Sales by Region"},
                                        "Create total Sales by Region chart from Orders")
            await call("create_chart", filepath=path, **routed)
            rejected = await tools["create_chart"].ainvoke({
                "filepath": path, "sheet_name": "Missing", "data_range": "D1:E3",
                "chart_type": "bar", "target_cell": "G2", "title": "Bad chart"})
            rejected_text = rejected[0]["text"] if isinstance(rejected, list) else str(rejected)
            self.assertTrue(rejected_text.startswith("Error:"), rejected)
            from agents.chart_review import verify_chart
            review = verify_chart(path, routed["sheet_name"], routed["data_range"], "Total Sales by Region")
            self.assertTrue(review["verified"])
            book = load_workbook(path)
            self.assertEqual(book["Orders"]["B4"].value, 8)
            self.assertEqual(len(book["Charts"]._charts), 1)
            self.assertEqual(book["_ChartData"].sheet_state, "hidden")
            self.assertFalse(any(cell.value is not None for cell in book["Charts"]._cells.values()))
            book.close()

    async def test_pivot_dashboard_and_pie_chart(self):
        from agents.graph import get_mcp_client, place_pivot_on_sheet
        from agents.chart_review import verify_chart

        tools = {tool.name: tool for tool in await get_mcp_client().get_tools()}
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "dashboard.xlsx")

            async def call(name, **args):
                result = await tools[name].ainvoke(args)
                self.assertFalse(str(result).lower().startswith("error:"), result)
                return result

            await call("create_workbook", filepath=path)
            await call("create_worksheet", filepath=path, sheet_name="Orders")
            await call("write_data_to_excel", filepath=path, sheet_name="Orders", start_cell="A1",
                       data=[["Category", "Profit"], ["Hardware", 4], ["Software", 7], ["Hardware", 6]])
            await call("create_pivot_table", filepath=path, sheet_name="Orders", data_range="A1:B4",
                       rows=["Category"], values=["Profit"], agg_func="sum")
            place_pivot_on_sheet(path, "Orders", "Dashboard", {"Sheet1", "Orders"})
            book = load_workbook(path)
            self.assertEqual(book["Dashboard"]["A1"].value, "Category")
            self.assertEqual(book["Dashboard"]["B1"].value, "Profit (sum)")
            self.assertEqual(book["Dashboard"]["B2"].value, 10)
            self.assertNotIn("Orders_pivot", book.sheetnames)
            book.close()
            from agents.chart_prepare import prepare_chart_call
            routed = prepare_chart_call(path, {"sheet_name": "Orders", "data_range": "A1:B4",
                                               "chart_type": "pie", "target_cell": "G2",
                                               "title": "Total Profit by Category"},
                                        "Create total Profit by Category pie chart from Orders")
            await call("create_chart", filepath=path, **routed)
            review = verify_chart(path, routed["sheet_name"], routed["data_range"],
                                  "Total Profit by Category", "pie")
            self.assertTrue(review["verified"])
            self.assertEqual(review["sheet"], "Charts")
            book = load_workbook(path)
            self.assertFalse(any(cell.value is not None for cell in book["Charts"]._cells.values()))
            self.assertEqual(book["_ChartData"].sheet_state, "hidden")
            book.close()


if __name__ == "__main__":
    unittest.main()
