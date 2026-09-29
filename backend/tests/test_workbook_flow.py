import os
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.worksheet.table import Table

from agents.chart_review import verify_chart
from workbook_sync import list_saved_files, updated_tables


def add_chart(sheet, title, first_col, first_row=1, last_row=3):
    chart = BarChart()
    chart.title = title
    chart.add_data(Reference(sheet, min_col=first_col + 1, min_row=first_row, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=first_col, min_row=first_row + 1, max_row=last_row))
    sheet.add_chart(chart, "D2")


class WorkbookFlowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.temp.name, "sales.xlsx")
        book = Workbook()
        orders = book.active
        orders.title = "Orders"
        orders.append(["Region", "Sales", "Segment"])
        orders.append(["Europe", 10, "Consumer"])
        orders.append(["Asia", 20, "Business"])
        orders.append(["Europe", 5, "Consumer"])
        charts = book.create_sheet("Charts")
        for row in [["Title", "Total Sales by Region"], ["Source", "Orders!A1:C4"],
                    ["Request", "SUM(Sales) by Region"], ["Type", "bar"]]:
            charts.append(row)
        for row, values in enumerate([["Region", "Sum of Sales"], ["Europe", 15], ["Asia", 20]], 1):
            charts.cell(row, 4, values[0])
            charts.cell(row, 5, values[1])
        add_chart(charts, "Total Sales by Region", 4)
        book.save(self.path)
        book.close()
        verify_chart(self.path, "Charts", "D1:E3", "Total Sales by Region")
        book = load_workbook(self.path)
        self.assertEqual(len(book["Charts"]._charts), 1)
        self.assertEqual(book["Orders"]["B4"].value, 5)
        book.close()

    def tearDown(self):
        self.temp.cleanup()

    def test_chart_saved_references_bounds_and_upsert(self):
        result = verify_chart(self.path, "Charts", "D1:E3", "Total Sales by Region")
        self.assertTrue(result["verified"])
        self.assertGreater(result["bounds"][0], 5)
        book = load_workbook(self.path)
        charts = book["Charts"]
        add_chart(charts, "Total Sales by Region", 4)
        book.save(self.path)
        book.close()

    def test_pie_chart_on_charts_sheet_avoids_dashboard_data(self):
        verify_chart(self.path, "Charts", "D1:E3", "Total Sales by Region")
        book = load_workbook(self.path)
        orders = book["Orders"]
        orders["D1"] = "Profit"
        orders["D2"], orders["D3"], orders["D4"] = 2, 4, 3
        dashboard = book.create_sheet("Dashboard")
        dashboard["A1"], dashboard["B1"] = "Category", "Sum of Profit"
        dashboard["A2"], dashboard["B2"] = "A", 6
        dashboard["A3"], dashboard["B3"] = "B", 3
        charts = book["Charts"]
        charts["B1"] = "Total Profit by Category"
        charts["B2"] = "Orders!A1:D4"
        charts["B3"] = "SUM(Profit) by Segment"
        charts["B4"] = "column"  # Descriptive metadata must not override the actual chart.
        charts["D1"], charts["E1"] = "Segment", "Sum of Profit"
        charts["D2"], charts["E2"] = "Consumer", 5
        charts["D3"], charts["E3"] = "Business", 4
        pie = PieChart()
        pie.title = "Total Profit by Category"
        pie.add_data(Reference(charts, min_col=5, min_row=1, max_row=3), titles_from_data=True)
        pie.set_categories(Reference(charts, min_col=4, min_row=2, max_row=3))
        charts.add_chart(pie, "D2")
        book.save(self.path)
        book.close()
        result = verify_chart(self.path, "Charts", "D1:E3", "Total Profit by Category", "pie")
        self.assertTrue(result["verified"])
        self.assertGreater(result["bounds"][0], 5)
        book = load_workbook(self.path)
        self.assertEqual(book["Dashboard"]["B2"].value, 6)
        self.assertEqual(len(book["Charts"]._charts), 2)
        book.close()

    def test_aggregate_source_recovers_from_summary_sheet_and_short_range(self):
        book = load_workbook(self.path)
        employees = book.create_sheet("Employees")
        employees.append(["Department", "Salary"])
        employees.append(["Engineering", 2])
        employees.append(["Sales", 7])
        employees.append(["Engineering", 3])
        summary = book.create_sheet("SalaryChart")
        summary.append(["Department", "Total Salary", None, "Title", "Total Salary by Department"])
        summary.append(["Engineering", 5, None, "Source", "SalaryChart!A1:B3"])
        summary.append(["Sales", 7, None, "Request", "Bar chart of total Salary by Department"])
        summary.cell(4, 4, "Type")
        summary.cell(4, 5, "bar")
        add_chart(summary, "Total Salary by Department", 1)
        book.save(self.path)
        book.close()

        review = verify_chart(self.path, "SalaryChart", "A1:B3", "Total Salary by Department", "bar")
        self.assertTrue(review["verified"])
        book = load_workbook(self.path)
        self.assertEqual(book["SalaryChart"]["E2"].value, "'Employees'!A1:B4")
        self.assertEqual(book["SalaryChart"]["E3"].value, "SUM(Salary) by Department")
        book["SalaryChart"]["E2"] = "Employees!A1:B3"
        book.save(self.path)
        book.close()
        verify_chart(self.path, "SalaryChart", "A1:B3", "Total Salary by Department", "bar")
        book = load_workbook(self.path)
        self.assertEqual(book["SalaryChart"]["E2"].value, "'Employees'!A1:B4")
        book.close()

    def test_second_chart_avoids_first(self):
        first = verify_chart(self.path, "Charts", "D1:E3", "Total Sales by Region")
        book = load_workbook(self.path)
        charts = book["Charts"]
        for row, values in enumerate([["Title", "Sales by Segment"], ["Source", "Orders!A1:C4"],
                                      ["Request", "Sales by Segment"], ["Type", "bar"]], 20):
            charts.cell(row, 1, values[0])
            charts.cell(row, 2, values[1])
        for row, values in enumerate([["Segment", "Sales"], ["Consumer", 15], ["Business", 20]], 20):
            charts.cell(row, 4, values[0])
            charts.cell(row, 5, values[1])
        add_chart(charts, "Sales by Segment", 4, 20, 22)
        book.save(self.path)
        book.close()
        second = verify_chart(self.path, "Charts", "D20:E22", "Sales by Segment")
        self.assertGreaterEqual(second["bounds"][1], first["bounds"][3])
        book = load_workbook(self.path)
        self.assertEqual(len(book["Charts"]._charts), 2)
        book.close()

    def test_same_title_from_different_source_keeps_both_charts(self):
        verify_chart(self.path, "Charts", "D1:E3", "Total Sales by Region")
        book = load_workbook(self.path)
        returns = book.create_sheet("Returns")
        returns.append(["Region", "Sales"])
        returns.append(["Europe", 7])
        returns.append(["Asia", 9])
        charts = book["Charts"]
        for row, values in enumerate([["Title", "Total Sales by Region"], ["Source", "Returns!A1:B3"],
                                      ["Request", "SUM(Sales) by Region"], ["Type", "bar"]], 20):
            charts.cell(row, 1, values[0])
            charts.cell(row, 2, values[1])
        for row, values in enumerate([["Region", "Sum of Sales"], ["Europe", 7], ["Asia", 9]], 20):
            charts.cell(row, 4, values[0])
            charts.cell(row, 5, values[1])
        add_chart(charts, "Total Sales by Region", 4, 20, 22)
        book.save(self.path)
        book.close()
        verify_chart(self.path, "Charts", "D20:E22", "Total Sales by Region")
        book = load_workbook(self.path)
        self.assertEqual(len(book["Charts"]._charts), 2)
        book.close()

    def test_incorrect_aggregate_is_rejected_without_rewrite(self):
        book = load_workbook(self.path)
        book["Charts"]["E2"] = 999
        book.save(self.path)
        book.close()
        with self.assertRaisesRegex(ValueError, "does not equal SUM"):
            verify_chart(self.path, "Charts", "D1:E3", "Total Sales by Region")
        self.assertEqual(load_workbook(self.path)["Charts"]["E2"].value, 999)

    def test_explicit_orders_destination_stays_outside_data(self):
        book = load_workbook(self.path)
        orders = book["Orders"]
        for row, values in enumerate([["Title", "Total Sales by Region"], ["Source", "Orders!A1:C4"],
                                      ["Request", "SUM(Sales) by Region"], ["Type", "bar"]], 1):
            orders.cell(row, 5, values[0])
            orders.cell(row, 6, values[1])
        for row, values in enumerate([["Region", "Sum of Sales"], ["Europe", 15], ["Asia", 20]], 1):
            orders.cell(row, 8, values[0])
            orders.cell(row, 9, values[1])
        add_chart(orders, "Total Sales by Region", 8)
        book.save(self.path)
        book.close()
        review = verify_chart(self.path, "Orders", "H1:I3", "Total Sales by Region")
        self.assertGreater(review["bounds"][0], 9)
        book = load_workbook(self.path)
        self.assertEqual(book["Orders"]["C4"].value, "Consumer")
        self.assertEqual(len(book["Charts"]._charts), 1)
        book.close()

    def test_wrong_series_references_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "series mismatch"):
            verify_chart(self.path, "Charts", "G1:H3", "Total Sales by Region")

    def test_falls_below_data_when_right_edge_is_full(self):
        book = load_workbook(self.path)
        book["Charts"].cell(1, 16380, "far right data")
        book.save(self.path)
        book.close()
        review = verify_chart(self.path, "Charts", "D1:E3", "Total Sales by Region")
        self.assertEqual(review["bounds"][0], 2)
        self.assertGreater(review["bounds"][1], 4)

    def test_inventory_and_updated_range(self):
        with open(os.path.join(self.temp.name, "broken.xlsx"), "wb") as stream:
            stream.write(b"not a workbook")
        self.assertEqual(list_saved_files(self.temp.name), ["sales.xlsx"])
        changes = updated_tables(self.temp.name, [{"name": "write_data_to_excel", "args": {
            "filepath": self.path, "sheet_name": "Orders", "start_cell": "B2", "data": [[12]]}}])
        self.assertEqual(changes[0]["workbook"], "sales.xlsx")
        self.assertEqual(changes[0]["headers"], ["Sales"])
        self.assertEqual(changes[0]["total_row_count"], 3)

    def test_named_excel_table_preview_uses_its_own_header_row(self):
        book = load_workbook(self.path)
        sheet = book["Orders"]
        sheet["E8"], sheet["F8"] = "Product", "Revenue"
        sheet["E9"], sheet["F9"] = "A", 30
        sheet["E10"], sheet["F10"] = "B", 40
        sheet.add_table(Table(displayName="RevenueTable", ref="E8:F10"))
        book.save(self.path)
        book.close()
        changes = updated_tables(self.temp.name, [{"name": "write_data_to_excel", "args": {
            "filepath": self.path, "sheet_name": "Orders", "start_cell": "F9", "data": [[30]]}}])
        self.assertEqual(changes[0]["range"], "F8:F10")
        self.assertEqual(changes[0]["headers"], ["Revenue"])
        self.assertEqual(changes[0]["rows"], [["30"], ["40"]])

    def test_multiple_changed_columns_share_a_short_preview(self):
        changes = updated_tables(self.temp.name, [
            {"name": "write_data_to_excel", "args": {"filepath": self.path, "sheet_name": "Orders",
                                                      "start_cell": "B2", "data": [[10]]}},
            {"name": "write_data_to_excel", "args": {"filepath": self.path, "sheet_name": "Orders",
                                                      "start_cell": "C2", "data": [["Consumer"]]}},
        ])
        self.assertEqual(changes[0]["headers"], ["Sales", "Segment"])
        self.assertEqual(changes[0]["range"], "B1:C4")


if __name__ == "__main__":
    unittest.main()
