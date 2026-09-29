
import os
import re
import tempfile
from typing import Optional

from fastapi import FastAPI, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel


from agents.graph import aask_agent
from workbook_sync import file_path as safe_file_path, list_saved_files, updated_tables, valid_file, workbook_directory
from fastapi.responses import FileResponse

# ─── App Setup ─────────────────────────────────────────────────────

app = FastAPI(title="AI Excel Analyst API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = workbook_directory()
os.makedirs(UPLOAD_DIR, exist_ok=True)


def concise_write_answer(result: dict, previews: list[dict], created_files: list[str]) -> str:
    """Report the saved outcome without echoing the agent's tool transcript."""
    operations = result.get("operations", [])
    errors = result.get("tool_errors", [])
    if errors:
        reason = str(errors[-1]).splitlines()[0].removeprefix("create_chart: ").removeprefix("Error: ")[:180]
        return f"I couldn't finish the request: {reason}"
    if not operations:
        return result.get("answer", "No response received.")
    charts = [op["args"] for op in operations if op["name"] == "create_chart"]
    pivots = [op["args"] for op in operations if op["name"] == "create_pivot_table"]
    if charts:
        chart = charts[-1]
        kind = chart.get("chart_type", "chart")
        title = chart.get("title") or "Chart"
        sheet = chart.get("sheet_name", "workbook")
        if pivots:
            pivot_sheet = pivots[-1].get("destination_sheet") or pivots[-1].get("sheet_name") or "the workbook"
            return f'Created the pivot table on {pivot_sheet} and "{title}" {kind} chart on {sheet}.'
        return f'Created "{title}" {kind} chart on {sheet}.'
    if created_files:
        return "Created " + ", ".join(created_files) + "."
    if pivots:
        sheet = pivots[-1].get("destination_sheet") or pivots[-1].get("sheet_name") or "the workbook"
        return f"Created the pivot table on {sheet}."
    if previews:
        preview = previews[-1]
        columns = ", ".join(str(item) for item in preview["headers"] if item)
        return f'Updated {preview["sheet"]}' + (f": {columns}." if columns else ".")
    sheets = [op["args"].get("sheet_name") for op in operations if op["name"] == "create_worksheet"]
    if sheets:
        return "Created " + ", ".join(sheets) + " sheet" + ("s." if len(sheets) > 1 else ".")
    return "Workbook updated."

# ─── Request Models ───────────────────────────────────────────────

class ChatRequest(BaseModel):
    question: str
    file_name: Optional[str] = None    # which uploaded file to analyze
    sheet_name: Optional[str] = None   # optional: specific sheet
    thread_id: Optional[str] = "default_session"
    session_id: Optional[str] = None

# ─── Endpoints ─────────────────────────────────────────────────────

@app.get("/")
def health():
    return {"status": "AI Excel Analyst running", "files_loaded": list_saved_files(UPLOAD_DIR)}


@app.post("/upload")
async def upload(file: UploadFile):
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename provided")

    # Validate file type
    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Only .xlsx and .xls files are supported")

    try:
        save_path = safe_file_path(UPLOAD_DIR, file.filename)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    content = await file.read()
    with tempfile.NamedTemporaryFile(dir=UPLOAD_DIR, suffix=os.path.splitext(save_path)[1], delete=False) as temporary:
        temporary.write(content)
        temporary_path = temporary.name
    try:
        if not valid_file(temporary_path):
            raise HTTPException(status_code=400, detail="Uploaded spreadsheet is invalid")
        os.replace(temporary_path, save_path)
    finally:
        if os.path.exists(temporary_path):
            os.remove(temporary_path)


    return {
        "file": file.filename,
        "path": save_path,
        "size_bytes": len(content),
        "message": f"File '{file.filename}' uploaded successfully",
    }


@app.get("/files")
def list_files():
    """List valid uploaded and agent-created files."""
    return {"files": list_saved_files(UPLOAD_DIR)}


@app.get("/download/{filename}")
async def download_file(filename: str):
    """Download an uploaded/modified file."""
    try:
        path = safe_file_path(UPLOAD_DIR, filename)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not os.path.exists(path):
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(path, filename=filename)


@app.get("/preview/{filename}")
async def preview_file(
    filename: str,
    sheet_name: Optional[str] = None,
    offset: int = 0,
    limit: int = 100,
    cell_range: Optional[str] = None,
):
    """Preview sheets and tabular data of an uploaded Excel file with pagination."""
    import pandas as pd
    try:
        path = safe_file_path(UPLOAD_DIR, filename)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not os.path.exists(path):
        raise HTTPException(status_code=404, detail="File not found")

    try:
        if cell_range and filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook
            from openpyxl.utils.cell import range_boundaries
            left, top, right, bottom = range_boundaries(cell_range)
            if not (1 <= left <= right <= 16384 and 1 <= top <= bottom <= 1048576):
                raise ValueError("Invalid preview range")
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                sheets = book.sheetnames
                target_sheet = sheet_name if sheet_name in sheets else sheets[0]
                values = list(book[target_sheet].iter_rows(min_row=top, max_row=bottom,
                                                           min_col=left, max_col=right, values_only=True))
            finally:
                book.close()
            headers = [str(value) if value is not None else "" for value in values[0]] if values else []
            df_full = pd.DataFrame(values[1:], columns=headers)
        else:
            with pd.ExcelFile(path) as xl:
                sheets = xl.sheet_names
                target_sheet = sheet_name if sheet_name and sheet_name in sheets else sheets[0]
                df_full = pd.read_excel(xl, sheet_name=target_sheet)

        total_rows = len(df_full)
        df_slice = df_full.iloc[offset : offset + limit].fillna("")

        return {
            "sheets": sheets,
            "active_sheet": target_sheet,
            "columns": list(df_slice.columns),
            "data": df_slice.to_dict(orient="records"),
            "total_rows": total_rows,
            "offset": offset,
            "limit": limit,
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to preview file: {str(e)}")



@app.post("/chat")
async def chat(data: ChatRequest):
    # Determine which file to analyze
    file_path = None
    files_before = list_saved_files(UPLOAD_DIR)
    if data.file_name:
        try:
            candidate = safe_file_path(UPLOAD_DIR, data.file_name)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        if data.file_name not in files_before:
            raise HTTPException(status_code=404, detail="Workbook not found")
        file_path = candidate
    elif files_before:
        file_path = safe_file_path(UPLOAD_DIR, files_before[-1])

    # Run the LangChain MCP agent with conversation memory
    thread_id = data.thread_id or data.session_id or "default_session"
    try:
        result = await aask_agent(
            question=data.question,
            file_path=file_path,
            sheet_name=data.sheet_name,
            thread_id=thread_id,
        )
        requested_chart = bool(re.search(r"\b(?:create|add|make|update|fix|move)\b[\s\S]*\bchart\b", data.question, re.I))
        if requested_chart and result.get("operations") and not any(
            op["name"] == "create_chart" for op in result["operations"]
        ) and not result.get("tool_errors"):
            result["tool_errors"] = ["create_chart: No verified chart was saved"]
        # These reads happen after execution and never cause the agent to rerun a write.
        files_after = list_saved_files(UPLOAD_DIR)
        result["files"] = files_after
        result["created_files"] = sorted(set(files_after) - set(files_before))
        chart_requested = any(op["name"] == "create_chart" for op in result.get("operations", [])) or any(
            str(error).startswith("create_chart:") for error in result.get("tool_errors", []))
        try:
            result["updated_tables"] = [] if chart_requested else updated_tables(UPLOAD_DIR, result.get("operations", []))
        except Exception as exc:
            result["updated_tables"] = []
            result["sync_error"] = f"Saved workbook changed, but its table preview could not load: {exc}"
        result["answer"] = concise_write_answer(result, result["updated_tables"], result["created_files"])
        result.pop("operations", None)
        result.pop("tool_errors", None)
        return result
    except Exception as e:
        files_after = list_saved_files(UPLOAD_DIR)
        return {
            "answer": f"Error during analysis: {str(e)}",
            "chart_image": None,
            "files": files_after,
            "created_files": sorted(set(files_after) - set(files_before)),
            "updated_tables": [],
        }


@app.delete("/files/{filename}")
async def delete_file(filename: str):
    """Remove an uploaded file."""
    if filename not in list_saved_files(UPLOAD_DIR):
        raise HTTPException(status_code=404, detail="File not found")
    path = safe_file_path(UPLOAD_DIR, filename)
    if os.path.exists(path):
        os.remove(path)
    return {"message": f"File '{filename}' deleted"}
