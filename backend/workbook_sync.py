"""Discover saved workbooks and build bounded previews after MCP writes."""

from __future__ import annotations

import os
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries, coordinate_from_string, column_index_from_string

MAX_PREVIEW_ROWS = 5
MAX_PREVIEW_COLUMNS = 30


def workbook_directory() -> str:
    """Resolve local storage beside the project, unless deployment overrides it."""
    return os.path.abspath(os.getenv("UPLOAD_DIR") or Path(__file__).resolve().parent.parent / "workbooks")


def file_path(directory: str, filename: str) -> str:
    if not filename or Path(filename).name != filename or filename.startswith("~"):
        raise ValueError("Invalid filename")
    if Path(filename).suffix.lower() not in {".xlsx", ".xls"}:
        raise ValueError("Unsupported file type")
    path = os.path.join(directory, filename)
    if os.path.normcase(os.path.dirname(os.path.realpath(path))) != os.path.normcase(os.path.realpath(directory)):
        raise ValueError("Workbook path escapes the shared directory")
    return path


def valid_file(path: str) -> bool:
    if not os.path.isfile(path):
        return False
    suffix = Path(path).suffix.lower()
    try:
        if suffix == ".xlsx":
            book = load_workbook(path, read_only=True)
            book.close()
        elif suffix == ".xls":
            import pandas as pd
            pd.ExcelFile(path).close()
        else:
            return False
        return True
    except Exception:
        return False


def list_saved_files(directory: str) -> list[str]:
    entries = sorted(
        (entry for entry in Path(directory).iterdir()
        if entry.is_file() and entry.resolve().parent == Path(directory).resolve()
        and entry.suffix.lower() in {".xlsx", ".xls"}
        and not entry.name.startswith("~$") and valid_file(str(entry))),
        key=lambda entry: (entry.stat().st_mtime_ns, entry.name.lower())
    )
    return [entry.name for entry in entries]


def _affected_bounds(args: dict):
    if args.get("data_range"):
        try:
            return range_boundaries(args["data_range"].split("!")[-1])
        except ValueError:
            pass
    if args.get("start_cell"):
        try:
            letters, row = coordinate_from_string(args["start_cell"])
            col = column_index_from_string(letters)
            data = args.get("data") or []
            height = len(data) if isinstance(data, list) else 1
            width = max((len(item) if isinstance(item, list) else 1 for item in data), default=1)
            return col, row, col + width - 1, row + height - 1
        except (ValueError, TypeError):
            pass
    if args.get("cell"):
        try:
            letters, row = coordinate_from_string(args["cell"])
            col = column_index_from_string(letters)
            return col, row, col, row
        except ValueError:
            pass
    return None


def updated_tables(directory: str, operations: list[dict]) -> list[dict]:
    """Read saved ranges once; a preview failure never replays a tool call."""
    targets = {}
    for operation in operations:
        name, args = operation["name"], operation["args"]
        if name not in {"write_data_to_excel", "apply_formula", "create_table", "copy_range", "insert_rows", "insert_columns", "delete_range"}:
            continue
        data = args.get("data") or []
        if name == "write_data_to_excel" and len(data) >= 4 and [str(row[0]).lower() for row in data[:4] if isinstance(row, list) and row] == ["title", "source", "request", "type"]:
            continue
        try:
            filename = Path(args["filepath"]).name
            if file_path(directory, filename) != args["filepath"]:
                continue
            sheet_name = args.get("sheet_name") or args.get("target_sheet")
            if not sheet_name:
                continue
            key = (filename, sheet_name)
            bounds = _affected_bounds(args)
            previous = targets.get(key)
            if bounds and previous:
                bounds = (min(bounds[0], previous[0]), min(bounds[1], previous[1]),
                          max(bounds[2], previous[2]), max(bounds[3], previous[3]))
            targets[key] = bounds or previous
        except (KeyError, ValueError):
            continue
    result = []
    for (filename, sheet_name), bounds in targets.items():
        path = file_path(directory, filename)
        if Path(path).suffix.lower() != ".xlsx" or not valid_file(path):
            continue
        book = load_workbook(path, read_only=False, data_only=True)
        try:
            if sheet_name not in book:
                continue
            sheet = book[sheet_name]
            # Use the table's header row but show only changed columns.
            for table in sheet.tables.values() if hasattr(sheet, "tables") else []:
                table_bounds = range_boundaries(table.ref)
                if bounds and table_bounds[0] <= bounds[0] <= table_bounds[2] and table_bounds[1] <= bounds[1] <= table_bounds[3]:
                    bounds = (bounds[0], table_bounds[1], bounds[2], table_bounds[3])
                    break
            if not bounds:
                bounds = (1, 1, min(sheet.max_column, MAX_PREVIEW_COLUMNS), sheet.max_row)
            left, top, right, bottom = bounds
            if top > 1 and sheet.cell(1, left).value is not None:
                top = 1
            if top == 1 and sheet.cell(1, left).value is not None:
                bottom = sheet.max_row
            right = min(max(right, left), left + MAX_PREVIEW_COLUMNS - 1)
            bottom = min(max(bottom, top), sheet.max_row)
            columns = [str(sheet.cell(top, col).value or get_column_letter(col)) for col in range(left, right + 1)]
            rows = [
                [str(sheet.cell(row, col).value) if sheet.cell(row, col).value is not None else ""
                 for col in range(left, right + 1)]
                for row in range(top + 1, min(bottom, top + MAX_PREVIEW_ROWS) + 1)
            ]
            result.append({"workbook": filename, "sheet": sheet_name,
                           "range": f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}",
                           "headers": columns, "rows": rows, "total_row_count": max(0, bottom - top)})
        finally:
            book.close()
    return result
